"""Raw-read label harmonization for DeepPrime and PRIDICT2.

The workbook label ``editing_efficiency_fraction`` is deliberately not used by
this module.  The functions below operate on exact read-class counts and create
the separately named ``harmonized_editing_efficiency_fraction``.
"""

from __future__ import annotations

import gzip
import math
import re
import statistics
from collections import defaultdict
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
import yaml
from openpyxl import load_workbook

from primetab.features.features import (
    CanonicalDesign,
    local_bases,
    normalize_dna,
    parse_edit_type,
    reconstruct_edited_target,
    stable_hash,
    validate_edit_alignment,
)
from primetab.paths import REPOSITORY_ROOT

HARMONIZATION_CONTRACT_PATH = (
    REPOSITORY_ROOT / "projects" / "prime-editing-tfm" / "configs" / "label_harmonization_v1.yaml"
)


def load_harmonization_contract() -> dict[str, Any]:
    """Load the frozen raw-read label definition."""

    with HARMONIZATION_CONTRACT_PATH.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def _dna(value: Any) -> str:
    sequence = str(value).strip().upper().replace("U", "T")
    if not sequence or set(sequence) - set("ACGT"):
        raise ValueError(f"invalid_dna_sequence:{value!r}")
    return sequence


def reverse_complement(sequence: str) -> str:
    return _dna(sequence).translate(str.maketrans("ACGT", "TGCA"))[::-1]


@dataclass(frozen=True)
class ReadClassCounts:
    """High-quality, exactly assigned target-site read counts."""

    n_wt: int
    n_intended: int
    n_other: int

    def __post_init__(self) -> None:
        for name, value in (
            ("n_wt", self.n_wt),
            ("n_intended", self.n_intended),
            ("n_other", self.n_other),
        ):
            if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                raise ValueError(f"invalid_read_count:{name}:{value!r}")

    @property
    def n_total(self) -> int:
        return self.n_wt + self.n_intended + self.n_other

    @property
    def intended_fraction(self) -> float | None:
        if self.n_total == 0:
            return None
        return self.n_intended / self.n_total

    @property
    def other_fraction(self) -> float | None:
        if self.n_total == 0:
            return None
        return self.n_other / self.n_total

    def __add__(self, other: object) -> ReadClassCounts:
        if not isinstance(other, ReadClassCounts):
            return NotImplemented
        return ReadClassCounts(
            n_wt=self.n_wt + other.n_wt,
            n_intended=self.n_intended + other.n_intended,
            n_other=self.n_other + other.n_other,
        )


@dataclass(frozen=True)
class CorrectedReplicate:
    """One biological replicate after common background correction."""

    raw_fraction: float | None
    training_eligible: bool
    failure_codes: tuple[str, ...]
    treated: ReadClassCounts
    control: ReadClassCounts
    background_intended_fraction: float | None


def background_corrected_fraction(
    treated: ReadClassCounts,
    control: ReadClassCounts,
    *,
    minimum_treated_depth: int = 200,
    minimum_control_depth: int = 200,
    maximum_background_intended_fraction: float = 0.05,
    maximum_control_other_fraction: float = 0.20,
) -> CorrectedReplicate:
    """Apply the frozen ``(p_treated - p_bg) / (1 - p_bg)`` definition.

    Values below zero or above one are retained in ``raw_fraction`` and fail
    training eligibility.  There is intentionally no clipping.
    """

    failures: list[str] = []
    if treated.n_total < minimum_treated_depth:
        failures.append("treated_depth_below_200")
    if control.n_total < minimum_control_depth:
        failures.append("control_depth_below_200")

    background = control.intended_fraction
    control_other = control.other_fraction
    if background is None or background >= 1:
        failures.append("invalid_background_denominator")
        return CorrectedReplicate(
            raw_fraction=None,
            training_eligible=False,
            failure_codes=tuple(dict.fromkeys(failures)),
            treated=treated,
            control=control,
            background_intended_fraction=background,
        )
    if background > maximum_background_intended_fraction:
        failures.append("background_intended_fraction_above_0_05")
    if control_other is not None and control_other > maximum_control_other_fraction:
        failures.append("control_other_fraction_above_0_20")

    observed = treated.intended_fraction
    raw = None if observed is None else (observed - background) / (1 - background)
    if raw is None or not math.isfinite(raw):
        failures.append("invalid_background_denominator")
    elif raw < 0:
        failures.append("corrected_fraction_below_0")
    elif raw > 1:
        failures.append("corrected_fraction_above_1")

    return CorrectedReplicate(
        raw_fraction=raw,
        training_eligible=not failures,
        failure_codes=tuple(dict.fromkeys(failures)),
        treated=treated,
        control=control,
        background_intended_fraction=background,
    )


def infer_edit_block(wt_sequence: str, intended_sequence: str) -> tuple[int, str, str]:
    """Infer the single contiguous replacement block between two alleles."""

    wt = _dna(wt_sequence)
    intended = _dna(intended_sequence)
    prefix = 0
    while prefix < min(len(wt), len(intended)) and wt[prefix] == intended[prefix]:
        prefix += 1
    suffix = 0
    max_suffix = min(len(wt) - prefix, len(intended) - prefix)
    while suffix < max_suffix and wt[-1 - suffix] == intended[-1 - suffix]:
        suffix += 1
    wt_end = len(wt) - suffix if suffix else len(wt)
    intended_end = len(intended) - suffix if suffix else len(intended)
    reference = wt[prefix:wt_end]
    alternate = intended[prefix:intended_end]
    if not reference and not alternate:
        raise ValueError("alleles_are_identical")
    return prefix, reference, alternate


def apply_edit_block(
    wide_wt_sequence: str,
    edit_index: int,
    reference: str,
    alternate: str,
) -> str:
    """Apply a previously verified edit block to a full target sequence."""

    wide = _dna(wide_wt_sequence)
    reference = _dna(reference) if reference else ""
    alternate = _dna(alternate) if alternate else ""
    if edit_index < 0 or edit_index + len(reference) > len(wide):
        raise ValueError("edit_block_out_of_bounds")
    if wide[edit_index : edit_index + len(reference)] != reference:
        raise ValueError("edit_block_reference_mismatch")
    return wide[:edit_index] + alternate + wide[edit_index + len(reference) :]


def common_quantification_windows(
    wide_wt_sequence: str,
    wide_intended_sequence: str,
    *,
    nick_index: int,
    rtt_length: int,
    downstream_flank: int = 2,
) -> tuple[str, str]:
    """Construct the shared nick-2 through RTT-end+2 allele windows.

    ``nick_index`` is the zero-based index of the first base after the primary
    nick in the WT sequence.  RTT length is measured in the edited flap, so the
    corresponding WT span is adjusted by the allele-length delta.
    """

    wt = _dna(wide_wt_sequence)
    intended = _dna(wide_intended_sequence)
    if nick_index < 2 or rtt_length <= 0 or downstream_flank < 0:
        raise ValueError("invalid_nick_or_rtt_length")
    delta = len(intended) - len(wt)
    start = nick_index - 2
    reference_flap_end = nick_index + rtt_length - delta
    wt_end = reference_flap_end + downstream_flank
    intended_end = reference_flap_end + delta + downstream_flank
    if wt_end > len(wt) or intended_end > len(intended):
        raise ValueError("quantification_window_out_of_bounds")
    wt_window = wt[start:wt_end]
    intended_window = intended[start:intended_end]
    if not wt_window or not intended_window:
        raise ValueError("empty_quantification_window")
    return wt_window, intended_window


def deepprime_common_windows(
    wide_wt_sequence: str,
    source_wt_window: str,
    source_intended_window: str,
    *,
    rtt_length: int,
    nick_index: int = 21,
    allow_additional_mismatches: bool = False,
) -> tuple[str, str]:
    """Reconstruct full DeepPrime alleles from Table S1 and slice v1 windows."""

    wide = _dna(wide_wt_sequence)
    source_wt = _dna(source_wt_window)
    source_intended = _dna(source_intended_window)
    source_start = wide.find(source_wt)
    if source_start < 0:
        raise ValueError("source_wt_window_not_in_wide_target")
    common_start = nick_index - 2
    source_offset = common_start - source_start
    if source_offset < 0:
        raise ValueError("source_window_does_not_cover_common_start")
    source_slices = (source_wt[source_offset:], source_intended[source_offset:])
    if allow_additional_mismatches:
        if not all(source_slices):
            raise ValueError("empty_source_common_slice")
        return source_slices
    local_index, reference, alternate = infer_edit_block(source_wt, source_intended)
    intended_wide = apply_edit_block(
        wide,
        source_start + local_index,
        reference,
        alternate,
    )
    calculated = common_quantification_windows(
        wide,
        intended_wide,
        nick_index=nick_index,
        rtt_length=rtt_length,
    )
    if source_slices != calculated:
        raise ValueError("source_window_common_slice_mismatch")
    return source_slices


def classify_allele(
    observed_window: str,
    wt_window: str,
    intended_window: str,
) -> str:
    """Classify one already extracted, high-quality quantification window."""

    observed = _dna(observed_window)
    wt = _dna(wt_window)
    intended = _dna(intended_window)
    if observed == wt:
        return "wt"
    if observed == intended:
        return "intended"
    return "other"


def _counts_from_frame(frame: pd.DataFrame) -> ReadClassCounts:
    return ReadClassCounts(
        n_wt=int(frame["n_wt"].sum()),
        n_intended=int(frame["n_intended"].sum()),
        n_other=int(frame["n_other"].sum()),
    )


def harmonize_count_frame(counts: pd.DataFrame) -> pd.DataFrame:
    """Create one harmonized label per observation from replicate-level counts.

    Required columns are ``observation_id``, ``design_id``,
    ``control_match_id``, ``control_group_id``, ``sample_role`` (``treated``
    or ``control``), ``biological_replicate``, ``technical_replicate``,
    ``n_wt``, ``n_intended``, and ``n_other``.  ``control_match_id`` identifies
    the concrete assayed library member or target context; unlike ``design_id``
    it must distinguish the same pegRNA design installed at different targets.
    Technical replicates are pooled by count; corrected biological-replicate
    fractions are averaged without read-depth weighting.
    """

    required = {
        "observation_id",
        "design_id",
        "control_match_id",
        "control_group_id",
        "sample_role",
        "biological_replicate",
        "technical_replicate",
        "n_wt",
        "n_intended",
        "n_other",
    }
    missing = sorted(required - set(counts.columns))
    if missing:
        raise ValueError(f"harmonization_count_columns_missing:{missing}")
    counts = counts.copy()
    key_columns = sorted(required - {"n_wt", "n_intended", "n_other"})
    if counts[key_columns].isna().any().any():
        null_columns = sorted(column for column in key_columns if counts[column].isna().any())
        raise ValueError(f"harmonization_count_keys_null:{null_columns}")
    if not set(counts["sample_role"].dropna().unique()) <= {"treated", "control"}:
        raise ValueError("invalid_sample_role")
    for column in ("n_wt", "n_intended", "n_other"):
        values = pd.to_numeric(counts[column], errors="coerce")
        if values.isna().any() or (values < 0).any() or (values % 1 != 0).any():
            raise ValueError(f"invalid_count_column:{column}")
        counts[column] = values.astype("int64")

    contract = load_harmonization_contract()
    qc = contract["training_qc"]
    replicate_policy = contract["replicate_aggregation"]
    minimum_biological_replicates = int(
        replicate_policy["minimum_eligible_biological_replicates"]
    )
    lineage_rollup_columns = set(replicate_policy.get("lineage_rollup_columns", ()))
    controls = counts.loc[counts["sample_role"].eq("control")]
    treated = counts.loc[counts["sample_role"].eq("treated")]
    control_counts_by_key = {
        key: _counts_from_frame(group)
        for key, group in controls.groupby(
            ["control_match_id", "control_group_id"],
            sort=False,
        )
    }
    metadata_excluded = required | {"sample_role", "biological_replicate", "technical_replicate"}
    metadata_columns = [column for column in counts.columns if column not in metadata_excluded]
    results: list[dict[str, Any]] = []

    for observation_id, observation_rows in treated.groupby("observation_id", sort=True):
        design_values = observation_rows["design_id"].dropna().unique()
        control_match_values = observation_rows["control_match_id"].dropna().unique()
        control_values = observation_rows["control_group_id"].dropna().unique()
        if (
            len(design_values) != 1
            or len(control_match_values) != 1
            or len(control_values) != 1
        ):
            raise ValueError(f"non_unique_observation_mapping:{observation_id}")
        design_id = design_values[0]
        control_match_id = control_match_values[0]
        control_group_id = control_values[0]
        observation_failures: list[str] = []
        control_counts = control_counts_by_key.get((control_match_id, control_group_id))
        if control_counts is None:
            observation_failures.append("missing_control_group")
            control_counts = ReadClassCounts(0, 0, 0)

        replicate_estimates: list[CorrectedReplicate] = []
        for _replicate, replicate_rows in observation_rows.groupby(
            "biological_replicate", sort=True
        ):
            replicate_estimates.append(
                background_corrected_fraction(
                    _counts_from_frame(replicate_rows),
                    control_counts,
                    minimum_treated_depth=int(qc["minimum_treated_depth_per_biological_replicate"]),
                    minimum_control_depth=int(qc["minimum_pooled_control_depth"]),
                    maximum_background_intended_fraction=float(
                        qc["maximum_background_intended_fraction"]
                    ),
                    maximum_control_other_fraction=float(qc["maximum_control_other_fraction"]),
                )
            )

        all_raw_values = [
            estimate.raw_fraction
            for estimate in replicate_estimates
            if estimate.raw_fraction is not None and math.isfinite(estimate.raw_fraction)
        ]
        eligible_values = [
            estimate.raw_fraction
            for estimate in replicate_estimates
            if estimate.training_eligible
            and estimate.raw_fraction is not None
            and math.isfinite(estimate.raw_fraction)
        ]
        replicate_failure_codes: list[str] = []
        for estimate in replicate_estimates:
            replicate_failure_codes.extend(estimate.failure_codes)
        if len(eligible_values) < minimum_biological_replicates:
            observation_failures.append("no_eligible_biological_replicate")

        raw_mean = statistics.fmean(all_raw_values) if all_raw_values else None
        eligible_mean = statistics.fmean(eligible_values) if eligible_values else None
        training_eligible = not observation_failures and eligible_mean is not None
        harmonized = eligible_mean if training_eligible else None
        pooled_treated = _counts_from_frame(observation_rows)
        pooled_estimate = background_corrected_fraction(
            pooled_treated,
            control_counts,
            minimum_treated_depth=0,
            minimum_control_depth=0,
            maximum_background_intended_fraction=1.0,
            maximum_control_other_fraction=1.0,
        )
        row: dict[str, Any] = {
            "observation_id": observation_id,
            "design_id": design_id,
            "control_match_id": control_match_id,
            "control_group_id": control_group_id,
            "harmonization_version": contract["harmonization_version"],
            "harmonized_editing_efficiency_fraction_raw": raw_mean,
            "harmonized_editing_efficiency_fraction": harmonized,
            "training_eligible": training_eligible,
            "harmonization_failure_codes": ";".join(dict.fromkeys(observation_failures)) or None,
            "excluded_replicate_failure_codes": (
                ";".join(dict.fromkeys(replicate_failure_codes)) or None
            ),
            "background_intended_fraction": control_counts.intended_fraction,
            "control_other_fraction": control_counts.other_fraction,
            "biological_replicate_count": len(replicate_estimates),
            "eligible_biological_replicate_count": len(eligible_values),
            "excluded_biological_replicate_count": (
                len(replicate_estimates) - len(eligible_values)
            ),
            "biological_replicate_sd": (
                statistics.stdev(eligible_values) if len(eligible_values) >= 2 else None
            ),
            "biological_replicate_sem": (
                statistics.stdev(eligible_values) / math.sqrt(len(eligible_values))
                if len(eligible_values) >= 2
                else None
            ),
            "pooled_count_sensitivity_fraction_raw": pooled_estimate.raw_fraction,
            "n_wt": pooled_treated.n_wt,
            "n_intended": pooled_treated.n_intended,
            "n_other": pooled_treated.n_other,
            "n_total": pooled_treated.n_total,
            "control_n_wt": control_counts.n_wt,
            "control_n_intended": control_counts.n_intended,
            "control_n_other": control_counts.n_other,
            "control_n_total": control_counts.n_total,
        }
        first = observation_rows.iloc[0]
        for column in metadata_columns:
            values = observation_rows[column].dropna().unique()
            if column in lineage_rollup_columns:
                row[column] = ";".join(sorted({str(value) for value in values})) or None
                continue
            if len(values) > 1:
                raise ValueError(f"non_unique_observation_metadata:{observation_id}:{column}")
            row[column] = values[0] if len(values) == 1 else first.get(column)
        results.append(row)

    return pd.DataFrame(results).sort_values("observation_id").reset_index(drop=True)


def _semicolon_values(value: Any) -> list[str]:
    return [part for part in str(value).split(";") if part and part.lower() != "nan"]


def audit_ena_run_manifest(
    path: Path,
    *,
    expected_bioproject: str,
    expected_run_count: int,
) -> dict[str, Any]:
    """Validate one ENA run manifest without downloading FASTQ files."""

    frame = pd.read_csv(path, sep="\t", dtype=str)
    required = {
        "run_accession",
        "study_accession",
        "library_layout",
        "read_count",
        "fastq_bytes",
        "fastq_md5",
        "fastq_ftp",
    }
    missing = sorted(required - set(frame.columns))
    if missing:
        raise ValueError(f"ena_manifest_columns_missing:{missing}")
    studies = sorted(frame["study_accession"].dropna().unique())
    run_count = len(frame)
    unique_runs = not frame["run_accession"].duplicated().any()
    paired_rows = frame["library_layout"].eq("PAIRED")
    two_fastqs = frame["fastq_ftp"].map(lambda value: len(_semicolon_values(value)) == 2)
    two_md5s = frame["fastq_md5"].map(lambda value: len(_semicolon_values(value)) == 2)
    two_sizes = frame["fastq_bytes"].map(lambda value: len(_semicolon_values(value)) == 2)
    total_fastq_bytes = sum(
        int(part) for value in frame["fastq_bytes"] for part in _semicolon_values(value)
    )
    total_reads = int(pd.to_numeric(frame["read_count"], errors="raise").sum())
    checks = {
        "bioproject_exact": studies == [expected_bioproject],
        "run_count_exact": run_count == expected_run_count,
        "run_accessions_unique": unique_runs,
        "all_runs_paired": bool(paired_rows.all()),
        "two_fastq_files_per_run": bool(two_fastqs.all()),
        "two_md5_values_per_run": bool(two_md5s.all()),
        "two_byte_counts_per_run": bool(two_sizes.all()),
    }
    return {
        "status": "pass" if all(checks.values()) else "fail",
        "path": str(path),
        "bioproject": studies,
        "run_count": run_count,
        "total_fastq_bytes": total_fastq_bytes,
        "total_read_pairs": total_reads,
        "checks": checks,
    }


@dataclass(frozen=True)
class FastqRecord:
    name: str
    sequence: str
    qualities: tuple[int, ...]


def _open_fastq(path: Path) -> BinaryIO:
    with path.open("rb") as probe:
        gzip_magic = probe.read(2) == b"\x1f\x8b"
    if gzip_magic:
        return gzip.open(path, "rb")
    return path.open("rb")


def iter_fastq(path: Path) -> Iterator[FastqRecord]:
    """Yield validated Phred+33 FASTQ records from plain or gzip input."""

    with _open_fastq(path) as handle:
        record_number = 0
        while True:
            header = handle.readline()
            if not header:
                break
            sequence = handle.readline()
            plus = handle.readline()
            quality = handle.readline()
            record_number += 1
            if not sequence or not plus or not quality:
                raise ValueError(f"truncated_fastq_record:{path}:{record_number}")
            if not header.startswith(b"@") or not plus.startswith(b"+"):
                raise ValueError(f"malformed_fastq_record:{path}:{record_number}")
            sequence_text = sequence.rstrip().decode("ascii").upper().replace("U", "T")
            quality_text = quality.rstrip().decode("ascii")
            if len(sequence_text) != len(quality_text):
                raise ValueError(f"fastq_sequence_quality_length:{path}:{record_number}")
            if set(sequence_text) - set("ACGTN"):
                raise ValueError(f"fastq_invalid_alphabet:{path}:{record_number}")
            yield FastqRecord(
                name=header[1:].decode("utf-8").split()[0].removesuffix("/1").removesuffix("/2"),
                sequence=sequence_text,
                qualities=tuple(ord(character) - 33 for character in quality_text),
            )


def iter_paired_fastq(read1: Path, read2: Path) -> Iterator[tuple[FastqRecord, FastqRecord]]:
    """Yield synchronized pairs and fail closed on unequal or mismatched files."""

    first = iter_fastq(read1)
    second = iter_fastq(read2)
    index = 0
    while True:
        left = next(first, None)
        right = next(second, None)
        if left is None and right is None:
            break
        index += 1
        if left is None or right is None:
            raise ValueError(f"unbalanced_fastq_pair:{index}")
        if left.name != right.name:
            raise ValueError(f"fastq_pair_name_mismatch:{index}:{left.name}:{right.name}")
        yield left, right


@dataclass(frozen=True)
class ExactReadDesign:
    """Source-layout adapter output consumed by the common exact classifier."""

    design_key: str
    primary_identity: str
    identity_tokens: tuple[str, ...]
    wt_window: str
    intended_window: str
    primary_read: int
    allele_read: int
    identity_allow_reverse_complement: bool = True
    allele_allow_reverse_complement: bool = False
    design_id: str | None = None
    target_locus_id: str | None = None
    observation_id: str | None = None
    control_match_id: str | None = None
    source_study: str | None = None
    source_sheet: str | None = None
    source_excel_row: int | None = None
    source_design_key: str | None = None
    source_wide_target: str | None = None
    source_declared_rtt_length: int | None = None
    normalized_rtt_length: int | None = None
    normalization_recovery: str | None = None
    preferred_identity_starts: tuple[int, ...] = ()

    def __post_init__(self) -> None:
        _dna(self.primary_identity)
        _dna(self.wt_window)
        _dna(self.intended_window)
        for token in self.identity_tokens:
            _dna(token)
        if self.primary_read not in (1, 2) or self.allele_read not in (1, 2):
            raise ValueError("read_number_must_be_1_or_2")
        if any(position < 0 for position in self.preferred_identity_starts):
            raise ValueError("identity_start_must_be_nonnegative")


def canonicalize_deepprime_table_s1_record(
    record: dict[str, Any],
    *,
    wide_column: str,
    allow_additional_rtt_mismatches: bool = False,
) -> CanonicalDesign:
    """Canonicalize one Table S1 member without fabricating a 74-nt flank.

    Table S1 contains a small number of 43-nt and 72-nt target records.  Every
    sequence component used by ``design_id`` is nevertheless present.  This
    adapter therefore validates the complete edit and constructs the canonical
    design while deliberately leaving target-locus construction to Table S2.
    """

    wide_target = normalize_dna(record[wide_column])
    pbs_length = int(record["PBS length"])
    source_rtt_length = int(record["RTT length"])
    extension = normalize_dna(record["3' extension sequence of pegRNA"])
    source_edit_type = str(record["Edit type"]).strip()
    if source_edit_type.lower() in {
        "transv",
        "transv2",
        "transi",
        "gtoc&ctoa&ctog",
    }:
        edit_type = "substitution"
    else:
        edit_type = parse_edit_type(source_edit_type)
    edit_position = int(record["Edit position"])
    edit_length = int(record["Edit length"])
    if len(extension) <= pbs_length:
        raise ValueError("table_s1_extension_not_longer_than_pbs")

    # The PBS is the invariant suffix of the transcript-oriented extension.
    # In the Lib-Off mismatch rows, the source ``RTT length`` is sometimes a
    # reference-coordinate length: it is one nucleotide longer for insertions
    # and one shorter for deletions than the physical RTT in the extension.
    # Splitting from the PBS suffix recovers the actual molecule and exactly
    # reproduces the RTlen values in DeepPrime Supplementary Table S2 sheet 20.
    rtt_length = len(extension) - pbs_length
    rtt_length_delta = rtt_length - source_rtt_length
    allowed_deltas = {0}
    if allow_additional_rtt_mismatches:
        if edit_type == "insertion":
            allowed_deltas.add(-edit_length)
        elif edit_type == "deletion":
            allowed_deltas.add(edit_length)
    if rtt_length_delta not in allowed_deltas:
        raise ValueError(
            "table_s1_rtt_length_semantics_mismatch:"
            f"source={source_rtt_length}:actual={rtt_length}:"
            f"edit_type={edit_type}:edit_length={edit_length}"
        )
    rtt = extension[:rtt_length]
    pbs = extension[rtt_length:]
    if len(wide_target) < 27:
        raise ValueError(f"table_s1_target_too_short:{len(wide_target)}")

    rtt_target = reverse_complement(rtt)
    edited_target, rtt_original_target = reconstruct_edited_target(
        wide_target,
        21,
        rtt_target,
        edit_type,
        edit_length,
    )
    validate_edit_alignment(
        rtt_original_target,
        rtt_target,
        edit_type,
        edit_position,
        edit_length,
        allow_additional_mismatches=allow_additional_rtt_mismatches,
    )
    right_homology_arm_length = (
        rtt_length - edit_position + 1
        if edit_type == "deletion"
        else rtt_length - edit_position - edit_length + 1
    )
    if right_homology_arm_length < 0:
        raise ValueError("negative_right_homology_arm")
    return CanonicalDesign(
        spacer=wide_target[4:24],
        pbs=pbs,
        rtt=rtt,
        pam_before=wide_target[24:27],
        pam_after=edited_target[24:27],
        edit_type=edit_type,
        edit_position=edit_position,
        edit_length=edit_length,
        right_homology_arm_length=right_homology_arm_length,
        rtt_original=reverse_complement(rtt_original_target),
        **local_bases(
            wide_target,
            21,
            edit_position,
            edit_type,
            edit_length,
        ),
    )


def load_deepprime_table_s1_designs(
    workbook_path: Path,
    *,
    sheet_name: str,
) -> list[ExactReadDesign]:
    """Build a DeepPrime exact-matching registry from Supplementary Table S1."""

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"deepprime_table_s1_sheet_missing:{sheet_name}")
        sheet = workbook[sheet_name]
        header_values = list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True)))
        last_header = max(index for index, value in enumerate(header_values) if value is not None)
        headers = header_values[: last_header + 1]
        wide_columns = [
            header
            for header in headers
            if isinstance(header, str) and header.startswith("Wide target sequence")
        ]
        if len(wide_columns) != 1:
            raise ValueError(f"deepprime_wide_target_header:{sheet_name}:{wide_columns}")
        wide_column = wide_columns[0]
        required = {
            "INDEX",
            "RTT length",
            "Extended Barcode",
            "Unedited (gRNA binding -2 ~ RTT binding +2)",
            "Prime-edited (gRNA binding -2 ~ RTT binding +2)",
            wide_column,
        }
        if not required <= set(headers):
            raise ValueError(
                f"deepprime_table_s1_columns_missing:{sheet_name}:{sorted(required - set(headers))}"
            )
        designs: list[ExactReadDesign] = []
        blank_streak = 0
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=5, max_col=len(headers), values_only=True),
            start=5,
        ):
            record = dict(zip(headers, values, strict=False))
            if record.get("INDEX") is None:
                blank_streak += 1
                if blank_streak >= 20:
                    break
                continue
            blank_streak = 0
            try:
                wt_window, intended_window = deepprime_common_windows(
                    record[wide_column],
                    record["Unedited (gRNA binding -2 ~ RTT binding +2)"],
                    record["Prime-edited (gRNA binding -2 ~ RTT binding +2)"],
                    rtt_length=int(record["RTT length"]),
                    allow_additional_mismatches=sheet_name == "Lib-Off",
                )
            except ValueError as exc:
                raise ValueError(
                    f"deepprime_window:{sheet_name}:{excel_row}:{record['INDEX']}:{exc}"
                ) from exc
            identity = _dna(record["Extended Barcode"])
            source_rtt_length = int(record["RTT length"])
            extension_length = len(normalize_dna(record["3' extension sequence of pegRNA"]))
            normalized_rtt_length = extension_length - int(record["PBS length"])
            try:
                canonical = canonicalize_deepprime_table_s1_record(
                    record,
                    wide_column=wide_column,
                    allow_additional_rtt_mismatches=sheet_name == "Lib-Off",
                )
            except ValueError as exc:
                raise ValueError(
                    f"deepprime_canonical:{sheet_name}:{excel_row}:"
                    f"{record['INDEX']}:{exc}"
                ) from exc
            wide_target = normalize_dna(record[wide_column])
            control_match = stable_hash(
                "control_match",
                {
                    "source_study": "DeepPrime",
                    "extended_barcode": identity,
                    "design_id": canonical.design_id,
                    "wide_target": wide_target,
                    "wt_window": wt_window,
                    "intended_window": intended_window,
                },
            )
            designs.append(
                ExactReadDesign(
                    design_key=stable_hash(
                        "read_design",
                        {
                            "source_study": "DeepPrime",
                            "source_sheet": sheet_name,
                            "source_excel_row": excel_row,
                            "source_design_key": str(record["INDEX"]),
                        },
                    ),
                    primary_identity=identity,
                    identity_tokens=(identity,),
                    wt_window=wt_window,
                    intended_window=intended_window,
                    primary_read=2,
                    allele_read=2,
                    identity_allow_reverse_complement=True,
                    allele_allow_reverse_complement=False,
                    design_id=canonical.design_id,
                    control_match_id=control_match,
                    source_study="DeepPrime",
                    source_sheet=sheet_name,
                    source_excel_row=excel_row,
                    source_design_key=str(record["INDEX"]),
                    source_wide_target=wide_target,
                    source_declared_rtt_length=source_rtt_length,
                    normalized_rtt_length=normalized_rtt_length,
                    normalization_recovery=(
                        "lib_off_reference_coordinate_rtt_length"
                        if normalized_rtt_length != source_rtt_length
                        else None
                    ),
                    preferred_identity_starts=(89, 90, 91),
                )
            )
        return designs
    finally:
        workbook.close()


PRIDICT2_TEVOPREQ1_DNA = "CGCGGTTCTATCTAGTTACGCGTTAAACCAACTAGAA"


def _location_start(value: Any) -> int:
    numbers = re.findall(r"-?\d+", str(value))
    if len(numbers) != 2:
        raise ValueError(f"invalid_location:{value!r}")
    start, end = (int(number) for number in numbers)
    if start < 0 or end <= start:
        raise ValueError(f"invalid_location_bounds:{value!r}")
    return start


def load_pridict2_sheet2_designs(workbook_path: Path) -> list[ExactReadDesign]:
    """Build the PRIDICT2 Library-Diverse exact-matching registry."""

    from primetab.features.workbooks import (
        canonicalize_pridict2_sheet2,
        pridict2_sheet2_wide_sequences,
    )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["2"]
        header_values = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        last_header = max(index for index, value in enumerate(header_values) if value is not None)
        headers = header_values[: last_header + 1]
        required = {
            "spacer",
            "Extension_seq",
            "barcode",
            "RTTlength",
            "RT_initial_location",
            "wide_initial_target",
            "wide_mutated_target",
        }
        if not required <= set(headers):
            raise ValueError(f"pridict2_sheet2_columns_missing:{sorted(required - set(headers))}")
        designs: list[ExactReadDesign] = []
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True),
            start=2,
        ):
            record = dict(zip(headers, values, strict=False))
            if record.get("spacer") is None:
                continue
            wide_wt, wide_intended, _recovery = pridict2_sheet2_wide_sequences(record)
            canonical, target_id, target_window = canonicalize_pridict2_sheet2(record)
            wt_window, intended_window = common_quantification_windows(
                wide_wt,
                wide_intended,
                nick_index=_location_start(record["RT_initial_location"]),
                rtt_length=int(record["RTTlength"]),
            )
            spacer = _dna(record["spacer"])
            extension = _dna(record["Extension_seq"])
            barcode = _dna(record["barcode"])
            control_match = stable_hash(
                "control_match",
                {
                    "source_study": "PRIDICT2",
                    "barcode": barcode,
                    "design_id": canonical.design_id,
                    "target_locus_id": target_id,
                    "wt_window": wt_window,
                    "intended_window": intended_window,
                },
            )
            designs.append(
                ExactReadDesign(
                    design_key=stable_hash(
                        "read_design",
                        {
                            "source_study": "PRIDICT2",
                            "source_sheet": "2",
                            "source_excel_row": excel_row,
                        },
                    ),
                    primary_identity=spacer,
                    identity_tokens=(spacer, extension, PRIDICT2_TEVOPREQ1_DNA, barcode),
                    wt_window=wt_window,
                    intended_window=intended_window,
                    primary_read=2,
                    allele_read=1,
                    identity_allow_reverse_complement=True,
                    allele_allow_reverse_complement=False,
                    design_id=canonical.design_id,
                    target_locus_id=target_id,
                    control_match_id=control_match,
                    source_study="PRIDICT2",
                    source_sheet="2",
                    source_excel_row=excel_row,
                    source_design_key=str(record.get("Name") or excel_row),
                    source_wide_target=target_window,
                )
            )
        return designs
    finally:
        workbook.close()


def _token_has_quality(
    record: FastqRecord,
    token: str,
    *,
    allow_reverse_complement: bool,
    minimum_phred: int,
) -> bool:
    candidates = (_dna(token),)
    if allow_reverse_complement:
        candidates += (reverse_complement(token),)
    for candidate in candidates:
        start = record.sequence.find(candidate)
        while start >= 0:
            if min(record.qualities[start : start + len(candidate)], default=-1) >= minimum_phred:
                return True
            start = record.sequence.find(candidate, start + 1)
    return False


def _token_is_present(
    record: FastqRecord,
    token: str,
    *,
    allow_reverse_complement: bool,
) -> bool:
    candidates = (_dna(token),)
    if allow_reverse_complement:
        candidates += (reverse_complement(token),)
    return any(candidate in record.sequence for candidate in candidates)


def count_exact_read_pairs(
    pairs: Iterable[tuple[FastqRecord, FastqRecord]],
    designs: Sequence[ExactReadDesign],
    *,
    maximum_pairs: int | None = None,
    minimum_mean_phred: float = 25,
    minimum_identity_phred: int = 20,
) -> tuple[pd.DataFrame, dict[str, int]]:
    """Reference exact matcher used for pilots and source-layout gold tests.

    Production adapters may replace candidate lookup with a faster compiled
    implementation, but must produce identical read classes on gold fixtures.
    """

    if len({design.design_key for design in designs}) != len(designs):
        raise ValueError("duplicate_exact_read_design_key")

    primary_indices: dict[tuple[int, int], dict[str, set[int]]] = defaultdict(
        lambda: defaultdict(set)
    )
    preferred_starts: dict[tuple[int, int], set[int]] = defaultdict(set)
    for index, design in enumerate(designs):
        token = _dna(design.primary_identity)
        index_key = (design.primary_read, len(token))
        lookup = primary_indices[index_key]
        lookup[token].add(index)
        if design.identity_allow_reverse_complement:
            lookup[reverse_complement(token)].add(index)
        preferred_starts[index_key].update(design.preferred_identity_starts)

    counts = {
        design.design_key: {
            "n_wt": 0,
            "n_intended": 0,
            "n_other": 0,
            "n_low_quality": 0,
        }
        for design in designs
    }
    audit = {
        "read_pairs": 0,
        "unassigned_identity": 0,
        "low_quality_identity": 0,
        "ambiguous_identity": 0,
        "low_quality": 0,
        "assigned": 0,
    }

    for left, right in pairs:
        if maximum_pairs is not None and audit["read_pairs"] >= maximum_pairs:
            break
        audit["read_pairs"] += 1
        candidate_ids: set[int] = set()
        for (read_number, length), index in primary_indices.items():
            record = left if read_number == 1 else right
            maximum_start = len(record.sequence) - length
            registered = sorted(
                start for start in preferred_starts[(read_number, length)] if start <= maximum_start
            )
            group_candidates: set[int] = set()
            for start in registered:
                group_candidates.update(index.get(record.sequence[start : start + length], ()))
            if group_candidates:
                candidate_ids.update(group_candidates)
                continue
            registered_set = set(registered)
            for start in range(maximum_start + 1):
                if start in registered_set:
                    continue
                group_candidates.update(index.get(record.sequence[start : start + length], ()))
            candidate_ids.update(group_candidates)
        verified: list[int] = []
        low_quality_identity = False
        for candidate_id in sorted(candidate_ids):
            design = designs[candidate_id]
            records = (left, right)
            has_quality = all(
                any(
                    _token_has_quality(
                        record,
                        token,
                        allow_reverse_complement=design.identity_allow_reverse_complement,
                        minimum_phred=minimum_identity_phred,
                    )
                    for record in records
                )
                for token in design.identity_tokens
            )
            if has_quality:
                verified.append(candidate_id)
                continue
            if all(
                any(
                    _token_is_present(
                        record,
                        token,
                        allow_reverse_complement=design.identity_allow_reverse_complement,
                    )
                    for record in records
                )
                for token in design.identity_tokens
            ):
                low_quality_identity = True
        if not verified:
            audit[
                "low_quality_identity" if low_quality_identity else "unassigned_identity"
            ] += 1
            continue
        if len(verified) != 1:
            audit["ambiguous_identity"] += 1
            continue

        design = designs[verified[0]]
        allele_record = left if design.allele_read == 1 else right
        mean_phred = statistics.fmean(allele_record.qualities)
        if mean_phred < minimum_mean_phred:
            counts[design.design_key]["n_low_quality"] += 1
            audit["low_quality"] += 1
            continue
        wt_candidates = (_dna(design.wt_window),)
        intended_candidates = (_dna(design.intended_window),)
        if design.allele_allow_reverse_complement:
            wt_candidates += (reverse_complement(design.wt_window),)
            intended_candidates += (reverse_complement(design.intended_window),)
        wt_match = any(token in allele_record.sequence for token in wt_candidates)
        intended_match = any(token in allele_record.sequence for token in intended_candidates)
        if wt_match and not intended_match:
            read_class = "n_wt"
        elif intended_match and not wt_match:
            read_class = "n_intended"
        else:
            read_class = "n_other"
        counts[design.design_key][read_class] += 1
        audit["assigned"] += 1

    rows = []
    design_lookup = {design.design_key: design for design in designs}
    for design_key in sorted(counts):
        design = design_lookup[design_key]
        row = {
            "design_key": design_key,
            "design_id": design.design_id,
            "target_locus_id": design.target_locus_id,
            "observation_id": design.observation_id,
            "control_match_id": design.control_match_id,
            "source_study": design.source_study,
            "source_sheet": design.source_sheet,
            "source_excel_row": design.source_excel_row,
            "source_design_key": design.source_design_key,
            **counts[design_key],
        }
        row["n_total"] = row["n_wt"] + row["n_intended"] + row["n_other"]
        rows.append(row)
    return pd.DataFrame(rows), audit
