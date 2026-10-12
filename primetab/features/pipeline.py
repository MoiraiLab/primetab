"""End-to-end audit, build, and verification for the unified v1 data package."""

from __future__ import annotations

import gc
import importlib.metadata
import json
import platform
import re
import shutil
import tempfile
from collections import Counter
try:
    from datetime import UTC, datetime
except ImportError:
    from datetime import datetime, timezone
    UTC = timezone.utc
from functools import lru_cache
from pathlib import Path
from typing import Any

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import yaml
from openpyxl import load_workbook

from primetab.assets import sha256
from primetab.features.features import (
    CanonicalDesign,
    compute_f0,
    compute_f1_sequence,
    stable_hash,
)
from primetab.features.harmonization import (
    HARMONIZATION_CONTRACT_PATH,
    load_harmonization_contract,
)
from primetab.features.workbooks import (
    SheetSpec,
    canonicalize_deepprime,
    canonicalize_pridict2_sheet2,
    canonicalize_pridict2_sheet7,
    condition_id,
    deepprime_condition,
    deepprime_specs,
    iter_sheet_records,
    load_deepprime_summary,
    observation_id,
    pridict2_sheet2_conditions,
    pridict2_sheet2_wide_sequences,
    pridict2_sheet7_conditions,
    pridict2_specs,
    raw_string,
    source_record_id,
    unique_headers,
)
from primetab.paths import REPOSITORY_ROOT, repository_relative

PROJECT_CONFIG_ROOT = REPOSITORY_ROOT / "projects" / "prime-editing-tfm" / "configs"
DATA_CONTRACT_PATH = PROJECT_CONFIG_ROOT / "data_contract_v2.yaml"
FEATURE_REGISTRY_PATH = PROJECT_CONFIG_ROOT / "feature_registry_v1.yaml"
DATA_DICTIONARY_PATH = PROJECT_CONFIG_ROOT / "data_dictionary_v2.yaml"
LABEL_ONTOLOGY_PATH = PROJECT_CONFIG_ROOT / "label_ontology_v1.yaml"

PUBLICATION_DEEPPRIME = "Yu_et_al_Cell_2023"
PUBLICATION_PRIDICT2 = "Mathis_et_al_NatBiotech_2024"


def _load_yaml(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def load_contract() -> dict[str, Any]:
    return _load_yaml(DATA_CONTRACT_PATH)


def load_feature_registry() -> dict[str, Any]:
    return _load_yaml(FEATURE_REGISTRY_PATH)


@lru_cache(maxsize=1)
def load_label_ontology() -> dict[str, Any]:
    """Load and structurally validate the frozen label ontology."""

    ontology = _load_yaml(LABEL_ONTOLOGY_PATH)
    definitions = ontology.get("definitions", {})
    mappings = ontology.get("source_mappings", [])
    expected_count = int(ontology["expected_source_mapping_count"])
    if len(mappings) != expected_count:
        raise ValueError(f"label_ontology_mapping_count:{len(mappings)}:{expected_count}")
    allowed_evidence = set(ontology.get("evidence_status_values", []))
    for definition_id, definition in definitions.items():
        if not definition.get("label_compatibility_group"):
            raise ValueError(f"label_definition_missing_group:{definition_id}")
        if definition.get("evidence_status") not in allowed_evidence:
            raise ValueError(f"label_definition_invalid_evidence_status:{definition_id}")
    seen: set[tuple[str, str, str]] = set()
    for mapping in mappings:
        key = (
            str(mapping["source_study"]),
            str(mapping["source_sheet"]),
            str(mapping["outcome_source_column"]),
        )
        if key in seen:
            raise ValueError(f"duplicate_label_ontology_mapping:{key}")
        seen.add(key)
        if mapping["label_definition_id"] not in definitions:
            raise ValueError(f"unknown_label_definition:{mapping['label_definition_id']}")
    return ontology


@lru_cache(maxsize=1)
def _label_ontology_lookup() -> dict[tuple[str, str, str], dict[str, str]]:
    ontology = load_label_ontology()
    definitions = ontology["definitions"]
    result: dict[tuple[str, str, str], dict[str, str]] = {}
    for mapping in ontology["source_mappings"]:
        definition_id = str(mapping["label_definition_id"])
        definition = definitions[definition_id]
        key = (
            str(mapping["source_study"]),
            str(mapping["source_sheet"]),
            str(mapping["outcome_source_column"]),
        )
        result[key] = {
            "label_ontology_version": str(ontology["ontology_version"]),
            "label_definition_id": definition_id,
            "label_compatibility_group": str(definition["label_compatibility_group"]),
            "label_evidence_status": str(definition["evidence_status"]),
        }
    return result


def label_registration(
    source_study: str,
    source_sheet: str,
    outcome_source_column: str,
) -> dict[str, str]:
    """Resolve one exact source tuple; unregistered labels are build errors."""

    key = (str(source_study), str(source_sheet), str(outcome_source_column))
    try:
        return _label_ontology_lookup()[key]
    except KeyError as exc:
        raise ValueError(f"unregistered_label_source_tuple:{key}") from exc


def feature_columns(registry: dict[str, Any], feature_set: str) -> list[str]:
    definition = registry["feature_sets"][feature_set]
    inherited = definition.get("inherits")
    columns = feature_columns(registry, inherited) if inherited else []
    return columns + list(definition["features"])


class ParquetBatchWriter:
    """Write dict records to a deterministic Parquet file in bounded batches."""

    def __init__(
        self,
        path: Path,
        schema: pa.Schema,
        batch_size: int = 10_000,
    ) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.schema = schema
        self.batch_size = batch_size
        self.buffer: list[dict[str, Any]] = []
        self.writer: pq.ParquetWriter | None = None
        self.rows = 0

    def write(self, record: dict[str, Any]) -> None:
        self.buffer.append(record)
        if len(self.buffer) >= self.batch_size:
            self.flush()

    def flush(self) -> None:
        if not self.buffer:
            return
        table = pa.Table.from_pylist(self.buffer, schema=self.schema)
        if self.writer is None:
            self.writer = pq.ParquetWriter(
                self.path,
                self.schema,
                compression="zstd",
                use_dictionary=True,
                write_statistics=True,
            )
        self.writer.write_table(table)
        self.rows += len(self.buffer)
        self.buffer.clear()

    def close(self) -> None:
        self.flush()
        if self.writer is not None:
            self.writer.close()
        else:
            pq.write_table(
                pa.Table.from_pylist([], schema=self.schema), self.path, compression="zstd"
            )


LINEAGE_SCHEMA = pa.schema(
    [
        ("source_record_id", pa.string()),
        ("source_study", pa.string()),
        ("source_sheet", pa.string()),
        ("source_excel_row", pa.int64()),
        ("record_role", pa.string()),
        ("publication_id", pa.string()),
    ]
)

QUARANTINE_SCHEMA = pa.schema(
    [
        ("source_record_id", pa.string()),
        ("source_study", pa.string()),
        ("source_sheet", pa.string()),
        ("source_excel_row", pa.int64()),
        ("failure_stage", pa.string()),
        ("failure_code", pa.string()),
        ("failure_detail", pa.string()),
    ]
)

ON_TARGET_SCHEMA = pa.schema(
    [
        ("observation_id", pa.string()),
        ("source_record_id", pa.string()),
        ("design_id", pa.string()),
        ("target_locus_id", pa.string()),
        ("condition_id", pa.string()),
        ("source_study", pa.string()),
        ("source_dataset_id", pa.string()),
        ("source_sheet", pa.string()),
        ("source_excel_row", pa.int64()),
        ("publication_id", pa.string()),
        ("outcome_type", pa.string()),
        ("outcome_source_column", pa.string()),
        ("outcome_source_value", pa.float64()),
        ("outcome_source_unit", pa.string()),
        ("editing_efficiency_fraction", pa.float64()),
        ("label_ontology_version", pa.string()),
        ("label_definition_id", pa.string()),
        ("label_compatibility_group", pa.string()),
        ("label_evidence_status", pa.string()),
        ("cell_line", pa.string()),
        ("editor_system", pa.string()),
        ("pam_family", pa.string()),
        ("scaffold_class", pa.string()),
        ("organism", pa.string()),
        ("mmr_status", pa.string()),
        ("assay_type", pa.string()),
        ("target_context", pa.string()),
        ("delivery_context", pa.string()),
        ("upstream_split", pa.string()),
        ("upstream_fold", pa.string()),
        ("upstream_model", pa.string()),
        ("normalization_status", pa.string()),
        ("failure_stage", pa.string()),
        ("failure_code", pa.string()),
    ]
)

MISMATCH_SCHEMA = pa.schema(
    [
        ("source_record_id", pa.string()),
        ("design_id", pa.string()),
        ("target_locus_id", pa.string()),
        ("source_group", pa.string()),
        ("source_target_group_id", pa.string()),
        ("linkage_status", pa.string()),
        ("linkage_failure_code", pa.string()),
        ("source_study", pa.string()),
        ("source_dataset_id", pa.string()),
        ("source_sheet", pa.string()),
        ("source_excel_row", pa.int64()),
        ("publication_id", pa.string()),
        ("condition_id", pa.string()),
        ("cell_line", pa.string()),
        ("editor_system", pa.string()),
        ("scaffold_class", pa.string()),
        ("target_number", pa.string()),
        ("source_index", pa.string()),
        ("pbs_length", pa.int64()),
        ("rtt_length", pa.int64()),
        ("mismatch_count", pa.int64()),
        ("mismatch_position", pa.string()),
        ("wild_to_mismatch", pa.string()),
        ("edit_position_source", pa.string()),
        ("edit_type_source", pa.string()),
        ("edit_length", pa.int64()),
        ("mismatch_type", pa.string()),
        ("average_on_target_efficiency_percent", pa.float64()),
        ("relative_complete_editing_efficiency", pa.float64()),
        ("relative_intended_only_editing_efficiency", pa.float64()),
        ("relative_mismatch_only_editing_efficiency", pa.float64()),
        ("relative_all_editing_efficiency", pa.float64()),
        ("normalization_status", pa.string()),
        ("failure_stage", pa.string()),
        ("failure_code", pa.string()),
    ]
)

TARGET_LOCUS_SCHEMA = pa.schema(
    [
        ("target_locus_id", pa.string()),
        ("oriented_74nt_window", pa.string()),
    ]
)

RECOVERY_SCHEMA = pa.schema(
    [
        ("source_record_id", pa.string()),
        ("source_study", pa.string()),
        ("source_sheet", pa.string()),
        ("source_excel_row", pa.int64()),
        ("recovery_code", pa.string()),
        ("evidence_fields", pa.string()),
    ]
)

DESIGN_SCHEMA = pa.schema(
    [
        ("design_id", pa.string()),
        ("spacer", pa.string()),
        ("pbs", pa.string()),
        ("rtt", pa.string()),
        ("extension", pa.string()),
        ("pam_before", pa.string()),
        ("pam_after", pa.string()),
        ("edit_type", pa.string()),
        ("edit_position", pa.int64()),
        ("edit_length", pa.int64()),
        ("right_homology_arm_length", pa.int64()),
        ("rtt_original", pa.string()),
        ("base_before_edit", pa.string()),
        ("base_after_edit", pa.string()),
        ("base_before_nick", pa.string()),
        ("base_after_nick", pa.string()),
    ]
)


def _float_or_none(value: Any) -> float | None:
    if value is None or value == "":
        return None
    return float(value)


def _int_or_none(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    integer = int(number)
    return integer if number == integer else None


def _outcome_fraction(value: Any) -> tuple[float | None, str | None]:
    if value is None or value == "":
        return None, None
    number = float(value)
    if not 0 <= number <= 100:
        return None, f"absolute_efficiency_out_of_range:{number}"
    return number / 100.0, None


def _source_declares_unknown_mismatch(record: dict[str, Any]) -> bool:
    """Recognize the workbook's complete unknown-mismatch sentinel tuple."""

    return (
        str(record.get("# of MM")).strip() == "?"
        and str(record.get("pos of MM")).strip() == "?"
        and str(record.get("Wild to MM")).strip() == "?"
        and str(record.get("type of MM")).strip() == "-"
    )


def _mismatch_source_target_group(record: dict[str, Any]) -> tuple[str, str]:
    """Return the source-defined target group available in mismatch sheets.

    DeepPrime mismatch sheets do not contain spacer, PBS, RTT, or target-window
    sequences, so a canonical design or target-locus hash cannot be derived.
    They do retain a stable ``GroupN`` prefix and ``Target #`` value.  The
    resulting identifier is valid only for leakage grouping inside the
    mismatch task and must never be treated as ``target_locus_id``.
    """

    source_index = raw_string(record.get("INDEX"))
    target_number = raw_string(record.get("Target #"))
    if not source_index or not target_number:
        raise ValueError("missing_source_target_group_fields")
    match = re.match(r"^(Group\d+)_", source_index)
    if match is None:
        raise ValueError(f"invalid_source_group:{source_index}")
    source_group = match.group(1)
    group_id = stable_hash(
        "source_target_group",
        {
            "source_study": "DeepPrime",
            "source_group": source_group,
            "target_number": target_number,
        },
    )
    return source_group, group_id


def _record_target_locus(target_loci: dict[str, str], locus_id: str, window: str) -> None:
    existing = target_loci.setdefault(locus_id, window)
    if existing != window:
        raise ValueError(f"target_locus_hash_collision:{locus_id}")


def _sheet_role(source: str, sheet: str) -> str:
    if source == "DeepPrime":
        if sheet == "Summary":
            return "summary_metadata"
        return "pegRNA_on_target" if int(sheet) <= 20 else "mismatch_off_target"
    if sheet == "Overview":
        return "overview_metadata"
    if sheet in {"2", "7"}:
        return "pegRNA_on_target"
    if sheet == "12":
        return "chromatin_context"
    return "supporting_metadata"


def _publication(source: str) -> str:
    return PUBLICATION_DEEPPRIME if source == "DeepPrime" else PUBLICATION_PRIDICT2


def _scan_source(
    source: str,
    path: Path,
    workbook_hash: str,
    specs: dict[str, SheetSpec],
    raw_root: Path,
    lineage: ParquetBatchWriter,
) -> tuple[list[dict[str, Any]], int]:
    data_workbook = load_workbook(path, read_only=True, data_only=True)
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    profiles: list[dict[str, Any]] = []
    total_rows = 0
    try:
        for sheet in data_workbook.sheetnames:
            spec = specs[sheet]
            worksheet = data_workbook[sheet]
            headers, records = iter_sheet_records(worksheet, spec, workbook_hash)
            raw_schema = pa.schema(
                [
                    ("__source_record_id", pa.string()),
                    ("__source_excel_row", pa.int64()),
                    *[(header, pa.string()) for header in headers],
                ]
            )
            raw_writer = ParquetBatchWriter(
                raw_root / f"{source.lower()}__{sheet}.parquet",
                raw_schema,
                batch_size=2_000,
            )
            type_counts: dict[str, Counter[str]] = {header: Counter() for header in headers}
            null_counts = Counter({header: 0 for header in headers})
            row_fingerprints: Counter[str] = Counter()
            error_cells = 0
            error_cells_by_column: Counter[str] = Counter()
            row_count = 0
            for excel_row, record_id, record, cells in records:
                raw_record: dict[str, Any] = {
                    "__source_record_id": record_id,
                    "__source_excel_row": excel_row,
                }
                for header, cell in zip(headers, cells, strict=True):
                    value = record[header]
                    if value is None:
                        null_counts[header] += 1
                    else:
                        type_counts[header][type(value).__name__] += 1
                    if cell.data_type == "e" or (isinstance(value, str) and value.startswith("#")):
                        error_cells += 1
                        error_cells_by_column[header] += 1
                    raw_record[header] = raw_string(value)
                row_fingerprints[
                    stable_hash("rawrow", [raw_record[header] for header in headers])
                ] += 1
                raw_writer.write(raw_record)
                role = _sheet_role(source, sheet)
                lineage.write(
                    {
                        "source_record_id": record_id,
                        "source_study": source,
                        "source_sheet": sheet,
                        "source_excel_row": excel_row,
                        "record_role": role,
                        "publication_id": _publication(source),
                    }
                )
                row_count += 1
            raw_writer.close()

            formula_count = 0
            formula_cached_missing = 0
            formula_cells_by_column: Counter[str] = Counter()
            formula_cached_missing_by_column: Counter[str] = Counter()
            formula_sheet = formula_workbook[sheet]
            last_row = (
                spec.first_data_row + spec.data_rows - 1
                if spec.data_rows is not None
                else formula_sheet.max_row
            )
            formula_rows = formula_sheet.iter_rows(
                min_row=spec.first_data_row, max_row=last_row, max_col=len(headers)
            )
            cached_rows = worksheet.iter_rows(
                min_row=spec.first_data_row, max_row=last_row, max_col=len(headers)
            )
            for formula_cells, cached_cells in zip(formula_rows, cached_rows, strict=True):
                for header, formula_cell, cached_cell in zip(
                    headers, formula_cells, cached_cells, strict=True
                ):
                    if formula_cell.data_type == "f":
                        formula_count += 1
                        formula_cells_by_column[header] += 1
                        formula_cached_missing += cached_cell.value is None
                        formula_cached_missing_by_column[header] += cached_cell.value is None

            profiles.append(
                {
                    "source_study": source,
                    "sheet": sheet,
                    "role": _sheet_role(source, sheet),
                    "columns": len(headers),
                    "rows": row_count,
                    "formula_cells": formula_count,
                    "formula_cells_with_cached_value": (formula_count - formula_cached_missing),
                    "formula_cells_missing_cached_value": formula_cached_missing,
                    "formula_cells_by_column": dict(formula_cells_by_column),
                    "formula_cells_missing_cached_value_by_column": dict(
                        formula_cached_missing_by_column
                    ),
                    "error_cells": error_cells,
                    "error_cells_by_column": dict(error_cells_by_column),
                    "exact_duplicate_rows": sum(
                        count - 1 for count in row_fingerprints.values() if count > 1
                    ),
                    "exact_duplicate_groups": sum(count > 1 for count in row_fingerprints.values()),
                    "null_counts": dict(null_counts),
                    "python_type_counts": {
                        key: dict(counter) for key, counter in type_counts.items()
                    },
                }
            )
            total_rows += row_count
    finally:
        data_workbook.close()
        formula_workbook.close()
    return profiles, total_rows


def audit_sources(deepprime_path: Path, pridict2_path: Path) -> dict[str, Any]:
    contract = load_contract()
    result: dict[str, Any] = {"status": "pass", "sources": {}, "checks": []}
    for source_key, path in (("deepprime", deepprime_path), ("pridict2", pridict2_path)):
        expected = contract["inputs"][source_key]
        actual_hash = sha256(path)
        actual_bytes = path.stat().st_size
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheetnames = workbook.sheetnames
        finally:
            workbook.close()
        source_result = {
            "path": repository_relative(path),
            "bytes": actual_bytes,
            "sha256": actual_hash,
            "sheet_count": len(sheetnames),
            "sheet_names": sheetnames,
        }
        result["sources"][source_key] = source_result
        checks = {
            "bytes": actual_bytes == int(expected["bytes"]),
            "sha256": actual_hash == expected["sha256"],
            "sheet_count": len(sheetnames) == int(expected["sheet_count"]),
        }
        for check, passed in checks.items():
            result["checks"].append({"source": source_key, "check": check, "passed": passed})
            if not passed:
                result["status"] = "fail"

    summary = load_deepprime_summary(deepprime_path)
    on_target_rows = sum(int(row["Total dataset size"]) for row in summary if row["Index"] <= 20)
    mismatch_rows = sum(int(row["Total dataset size"]) for row in summary if row["Index"] >= 21)
    deepprime_counts = {
        "summary_datasets": len(summary),
        "total_rows": on_target_rows + mismatch_rows,
        "on_target_rows": on_target_rows,
        "mismatch_rows": mismatch_rows,
    }
    result["sources"]["deepprime"]["declared_counts"] = deepprime_counts
    expected_dp = contract["inputs"]["deepprime"]["expected_rows"]
    for key, expected_key in (
        ("total_rows", "total"),
        ("on_target_rows", "on_target_sheets_1_20"),
        ("mismatch_rows", "mismatch_sheets_21_30"),
    ):
        passed = deepprime_counts[key] == int(expected_dp[expected_key])
        result["checks"].append({"source": "deepprime", "check": key, "passed": passed})
        if not passed:
            result["status"] = "fail"

    pridict_workbook = load_workbook(pridict2_path, read_only=True, data_only=True)
    try:
        sheet2 = pridict_workbook["2"]
        label_columns = pridict2_sheet2_conditions()
        headers2 = unique_headers(
            cell.value for cell in next(sheet2.iter_rows(min_row=1, max_row=1))
        )
        label_indexes = {name: headers2.index(name) for name in label_columns}
        sheet2_designs = sheet2_observations = 0
        for values in sheet2.iter_rows(min_row=2, max_col=len(headers2), values_only=True):
            if all(value is None for value in values):
                continue
            sheet2_designs += 1
            sheet2_observations += sum(
                values[index] is not None for index in label_indexes.values()
            )

        sheet7 = pridict_workbook["7"]
        label_columns7 = pridict2_sheet7_conditions()
        headers7 = unique_headers(
            cell.value for cell in next(sheet7.iter_rows(min_row=1, max_row=1))
        )
        label_indexes7 = {name: headers7.index(name) for name in label_columns7}
        sheet7_designs = sheet7_observations = 0
        for values in sheet7.iter_rows(min_row=2, max_col=len(headers7), values_only=True):
            if all(value is None for value in values):
                continue
            sheet7_designs += 1
            sheet7_observations += sum(
                values[index] is not None for index in label_indexes7.values()
            )

        sheet12 = pridict_workbook["12"]
        headers12 = unique_headers(
            cell.value for cell in next(sheet12.iter_rows(min_row=1, max_row=1))
        )
        pe_index = headers12.index("PE_editing_efficiency")
        sheet12_rows = sheet12_labels = 0
        for values in sheet12.iter_rows(min_row=2, max_col=len(headers12), values_only=True):
            if all(value is None for value in values):
                continue
            sheet12_rows += 1
            sheet12_labels += values[pe_index] is not None
    finally:
        pridict_workbook.close()

    pridict_counts = {
        "sheet_2_designs": sheet2_designs,
        "sheet_2_observations": sheet2_observations,
        "sheet_7_designs": sheet7_designs,
        "sheet_7_observations": sheet7_observations,
        "sheet_12_rows": sheet12_rows,
        "sheet_12_pe_labels": sheet12_labels,
    }
    result["sources"]["pridict2"]["observed_counts"] = pridict_counts
    expected_pridict = contract["inputs"]["pridict2"]["expected_rows"]
    for key, value in pridict_counts.items():
        passed = value == int(expected_pridict[key])
        result["checks"].append({"source": "pridict2", "check": key, "passed": passed})
        if not passed:
            result["status"] = "fail"
    if result["status"] != "pass":
        raise ValueError(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def _validate_known_source_anomalies(
    profiles: list[dict[str, Any]], contract: dict[str, Any]
) -> dict[str, Any]:
    """Lock immutable upstream anomalies to explicit, reviewed expectations."""

    profile_by_key = {
        (str(profile["source_study"]), str(profile["sheet"])): profile for profile in profiles
    }
    registered_error_keys: set[tuple[str, str]] = set()
    registered_duplicate_keys: set[tuple[str, str]] = set()
    checks: list[dict[str, Any]] = []
    for anomaly in contract.get("known_source_anomalies", []):
        source = str(anomaly["source_study"])
        sheet = str(anomaly["sheet"])
        key = (source, sheet)
        profile = profile_by_key.get(key)
        if profile is None:
            raise ValueError(f"known_source_anomaly_sheet_missing:{source}:{sheet}")
        anomaly_type = str(anomaly["anomaly_type"])
        if anomaly_type == "formula_error_cells":
            registered_error_keys.add(key)
            column = str(anomaly["source_column"])
            observed = {
                "error_cells": int(profile["error_cells_by_column"].get(column, 0)),
                "formula_cells": int(profile["formula_cells_by_column"].get(column, 0)),
                "formula_cells_missing_cached_value": int(
                    profile["formula_cells_missing_cached_value_by_column"].get(column, 0)
                ),
            }
            expected = {
                "error_cells": int(anomaly["expected_error_cells"]),
                "formula_cells": int(anomaly["expected_formula_cells"]),
                "formula_cells_missing_cached_value": int(
                    anomaly["expected_formula_cells_missing_cached_value"]
                ),
            }
        elif anomaly_type == "exact_duplicate_rows":
            registered_duplicate_keys.add(key)
            observed = {
                "exact_duplicate_rows": int(profile["exact_duplicate_rows"]),
                "exact_duplicate_groups": int(profile["exact_duplicate_groups"]),
            }
            expected = {
                "exact_duplicate_rows": int(anomaly["expected_duplicate_rows"]),
                "exact_duplicate_groups": int(anomaly["expected_duplicate_groups"]),
            }
        else:
            raise ValueError(f"unsupported_known_source_anomaly:{anomaly_type}")
        passed = observed == expected
        checks.append(
            {
                "anomaly_id": anomaly["id"],
                "source_study": source,
                "sheet": sheet,
                "anomaly_type": anomaly_type,
                "observed": observed,
                "expected": expected,
                "passed": passed,
                "downstream_policy": anomaly["downstream_policy"],
            }
        )
        if not passed:
            raise ValueError(f"known_source_anomaly_drift:{anomaly['id']}:{observed}:{expected}")

    unregistered: list[dict[str, Any]] = []
    for profile in profiles:
        key = (str(profile["source_study"]), str(profile["sheet"]))
        if int(profile["error_cells"]) and key not in registered_error_keys:
            unregistered.append(
                {
                    "source_study": key[0],
                    "sheet": key[1],
                    "anomaly_type": "formula_error_cells",
                    "rows": int(profile["error_cells"]),
                }
            )
        if int(profile["exact_duplicate_rows"]) and key not in registered_duplicate_keys:
            unregistered.append(
                {
                    "source_study": key[0],
                    "sheet": key[1],
                    "anomaly_type": "exact_duplicate_rows",
                    "rows": int(profile["exact_duplicate_rows"]),
                }
            )
    if unregistered:
        raise ValueError(f"unregistered_source_anomalies:{unregistered}")
    return {
        "status": "pass",
        "registered_anomalies": checks,
        "unregistered_anomalies": unregistered,
    }


def _quarantine(
    writer: ParquetBatchWriter,
    source: str,
    sheet: str,
    excel_row: int,
    record_id: str,
    stage: str,
    error: Exception | str,
) -> None:
    detail = str(error)
    if isinstance(error, Exception):
        code = detail.split(":", 1)[0] or error.__class__.__name__
    else:
        code = detail.split(":", 1)[0] or "unspecified_failure"
    writer.write(
        {
            "source_record_id": record_id,
            "source_study": source,
            "source_sheet": sheet,
            "source_excel_row": excel_row,
            "failure_stage": stage,
            "failure_code": code,
            "failure_detail": detail[:2_000],
        }
    )


def _observation_record(
    *,
    record_id: str,
    design_id: str | None,
    locus_id: str | None,
    condition: dict[str, str],
    source_study: str,
    source_dataset_id: str,
    source_sheet: str,
    source_excel_row: int,
    outcome_column: str,
    outcome_value: Any,
    upstream_split: Any = None,
    upstream_fold: Any = None,
    upstream_model: Any = None,
    normalization_status: str = "ok",
    failure_stage: str | None = None,
    failure_code: str | None = None,
) -> tuple[dict[str, Any], str | None]:
    fraction, label_error = _outcome_fraction(outcome_value)
    registration = label_registration(source_study, source_sheet, outcome_column)
    status = normalization_status if label_error is None else "quarantined"
    stages = [value for value in (failure_stage, "label" if label_error else None) if value]
    codes = [
        value
        for value in (
            failure_code,
            label_error.split(":", 1)[0] if label_error else None,
        )
        if value
    ]
    return (
        {
            "observation_id": observation_id(record_id, outcome_column),
            "source_record_id": record_id,
            "design_id": design_id,
            "target_locus_id": locus_id,
            "condition_id": condition_id(condition),
            "source_study": source_study,
            "source_dataset_id": source_dataset_id,
            "source_sheet": source_sheet,
            "source_excel_row": source_excel_row,
            "publication_id": _publication(source_study),
            "outcome_type": "absolute_editing_efficiency",
            "outcome_source_column": outcome_column,
            "outcome_source_value": _float_or_none(outcome_value),
            "outcome_source_unit": "percent",
            "editing_efficiency_fraction": fraction,
            **registration,
            **condition,
            "upstream_split": raw_string(upstream_split),
            "upstream_fold": raw_string(upstream_fold),
            "upstream_model": raw_string(upstream_model),
            "normalization_status": status,
            "failure_stage": "|".join(stages) if stages else None,
            "failure_code": "|".join(codes) if codes else None,
        },
        label_error,
    )


def _process_deepprime(
    path: Path,
    workbook_hash: str,
    summary: list[dict[str, Any]],
    on_target: ParquetBatchWriter,
    mismatch: ParquetBatchWriter,
    quarantine: ParquetBatchWriter,
    designs: dict[str, CanonicalDesign],
    target_loci: dict[str, str],
) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    by_index = {int(row["Index"]): row for row in summary}
    counts = Counter()
    source_feature_differences = Counter()
    mismatch_source_target_groups: set[str] = set()
    try:
        for sheet_index in range(1, 31):
            sheet = str(sheet_index)
            summary_row = by_index[sheet_index]
            spec = SheetSpec(4, 5, int(summary_row["Total dataset size"]))
            _headers, records = iter_sheet_records(workbook[sheet], spec, workbook_hash)
            condition = deepprime_condition(summary_row)
            if sheet_index <= 20:
                for excel_row, record_id, record, _cells in records:
                    design: CanonicalDesign | None = None
                    locus_id: str | None = None
                    status = "ok"
                    failure_stage = None
                    failure_code = None
                    try:
                        try:
                            design, locus_id, _window = canonicalize_deepprime(record)
                        except ValueError as exc:
                            if sheet_index != 20 or str(exc) != "edit_alignment_mismatch":
                                raise
                            design, locus_id, _window = canonicalize_deepprime(
                                record,
                                allow_additional_rtt_mismatches=True,
                            )
                            source_feature_differences[
                                "library_off_rows_with_additional_rtt_mismatch"
                            ] += 1
                        designs.setdefault(design.design_id, design)
                        _record_target_locus(target_loci, locus_id, _window)
                        source_feature_differences["source_rha_mismatch"] += int(
                            int(float(record["RHA_len"])) != design.right_homology_arm_length
                        )
                    except (KeyError, TypeError, ValueError) as exc:
                        status = "quarantined"
                        failure_stage = "canonicalization"
                        failure_code = str(exc).split(":", 1)[0]
                        _quarantine(
                            quarantine,
                            "DeepPrime",
                            sheet,
                            excel_row,
                            record_id,
                            "canonicalization",
                            exc,
                        )
                    observation, label_error = _observation_record(
                        record_id=record_id,
                        design_id=design.design_id if design else None,
                        locus_id=locus_id,
                        condition=condition,
                        source_study="DeepPrime",
                        source_dataset_id=str(summary_row["ID"]),
                        source_sheet=sheet,
                        source_excel_row=excel_row,
                        outcome_column="Measured_PE_efficiency",
                        outcome_value=record.get("Measured_PE_efficiency"),
                        upstream_split=record.get("Data set name"),
                        upstream_fold=record.get("Fold"),
                        upstream_model=record.get("Trained model"),
                        normalization_status=status,
                        failure_stage=failure_stage,
                        failure_code=failure_code,
                    )
                    on_target.write(observation)
                    if label_error:
                        _quarantine(
                            quarantine,
                            "DeepPrime",
                            sheet,
                            excel_row,
                            record_id,
                            "label",
                            label_error,
                        )
                    counts["on_target"] += 1
            else:
                for excel_row, record_id, record, _cells in records:
                    source_unknown_mismatch = _source_declares_unknown_mismatch(record)
                    source_group: str | None = None
                    source_target_group_id: str | None = None
                    try:
                        source_group, source_target_group_id = _mismatch_source_target_group(record)
                        mismatch_source_target_groups.add(source_target_group_id)
                    except ValueError as exc:
                        _quarantine(
                            quarantine,
                            "DeepPrime",
                            sheet,
                            excel_row,
                            record_id,
                            "mismatch_linkage",
                            exc,
                        )
                    parsed_integer_fields = {
                        "pbs_length": _int_or_none(record.get("PBS len")),
                        "rtt_length": _int_or_none(record.get("RTT len")),
                        "mismatch_count": (
                            None if source_unknown_mismatch else _int_or_none(record.get("# of MM"))
                        ),
                        "edit_length": _int_or_none(record.get("Edit length")),
                    }
                    invalid_integer_fields = [
                        name
                        for name, value in parsed_integer_fields.items()
                        if value is None
                        and not (name == "mismatch_count" and source_unknown_mismatch)
                    ]
                    if invalid_integer_fields:
                        _quarantine(
                            quarantine,
                            "DeepPrime",
                            sheet,
                            excel_row,
                            record_id,
                            "mismatch_field_parse",
                            "invalid_integer_fields:" + ",".join(invalid_integer_fields),
                        )
                    if source_unknown_mismatch:
                        counts["source_declared_unknown_mismatch_rows"] += 1
                    mismatch_status = (
                        "quarantined"
                        if invalid_integer_fields or source_target_group_id is None
                        else (
                            "complete_with_source_unknown_fields"
                            if source_unknown_mismatch
                            else "ok"
                        )
                    )
                    mismatch.write(
                        {
                            "source_record_id": record_id,
                            "design_id": None,
                            "target_locus_id": None,
                            "source_group": source_group,
                            "source_target_group_id": source_target_group_id,
                            "linkage_status": (
                                "source_group_only"
                                if source_target_group_id is not None
                                else "unavailable"
                            ),
                            "linkage_failure_code": (
                                "canonical_sequence_fields_absent_from_source_sheet"
                                if source_target_group_id is not None
                                else "source_target_group_unavailable"
                            ),
                            "source_study": "DeepPrime",
                            "source_dataset_id": str(summary_row["ID"]),
                            "source_sheet": sheet,
                            "source_excel_row": excel_row,
                            "publication_id": PUBLICATION_DEEPPRIME,
                            "condition_id": condition_id(condition),
                            "cell_line": condition["cell_line"],
                            "editor_system": condition["editor_system"],
                            "scaffold_class": condition["scaffold_class"],
                            "target_number": raw_string(record.get("Target #")),
                            "source_index": raw_string(record.get("INDEX")),
                            "pbs_length": parsed_integer_fields["pbs_length"],
                            "rtt_length": parsed_integer_fields["rtt_length"],
                            "mismatch_count": parsed_integer_fields["mismatch_count"],
                            "mismatch_position": raw_string(record.get("pos of MM")),
                            "wild_to_mismatch": raw_string(record.get("Wild to MM")),
                            "edit_position_source": raw_string(record.get("Edit pos")),
                            "edit_type_source": raw_string(record.get("Edit type")),
                            "edit_length": parsed_integer_fields["edit_length"],
                            "mismatch_type": raw_string(record.get("type of MM")),
                            "average_on_target_efficiency_percent": _float_or_none(
                                record.get("Avg.On-target efficiency")
                            ),
                            "relative_complete_editing_efficiency": _float_or_none(
                                record.get("Relative complete editing efficiency")
                            ),
                            "relative_intended_only_editing_efficiency": _float_or_none(
                                record.get("Relative intended only editing efficiency")
                            ),
                            "relative_mismatch_only_editing_efficiency": _float_or_none(
                                record.get("Relative mismatch only editing efficiency")
                            ),
                            "relative_all_editing_efficiency": _float_or_none(
                                record.get("Relative all editing efficiency")
                            ),
                            "normalization_status": mismatch_status,
                            "failure_stage": (
                                "mismatch_field_parse"
                                if invalid_integer_fields
                                else (
                                    "mismatch_linkage"
                                    if source_target_group_id is None
                                    else ("source_missingness" if source_unknown_mismatch else None)
                                )
                            ),
                            "failure_code": (
                                "invalid_integer_fields"
                                if invalid_integer_fields
                                else (
                                    "source_target_group_unavailable"
                                    if source_target_group_id is None
                                    else (
                                        "source_declared_unknown_mismatch_metadata"
                                        if source_unknown_mismatch
                                        else None
                                    )
                                )
                            ),
                        }
                    )
                    counts["mismatch"] += 1
    finally:
        workbook.close()
    return {
        "counts": dict(counts),
        "source_recalculation_checks": dict(source_feature_differences),
        "mismatch_source_target_groups": len(mismatch_source_target_groups),
        "mismatch_linkage_policy": (
            "canonical design_id and target_locus_id remain null because source sheets "
            "contain no sequence fields; source_target_group_id is the only supported "
            "within-task leakage group"
        ),
    }


def _process_pridict2(
    path: Path,
    workbook_hash: str,
    on_target: ParquetBatchWriter,
    quarantine: ParquetBatchWriter,
    designs: dict[str, CanonicalDesign],
    target_loci: dict[str, str],
    recoveries: ParquetBatchWriter,
) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    counts = Counter()
    try:
        conditions2 = pridict2_sheet2_conditions()
        _headers, records = iter_sheet_records(workbook["2"], SheetSpec(1, 2), workbook_hash)
        for excel_row, record_id, record, _cells in records:
            design: CanonicalDesign | None = None
            locus_id: str | None = None
            status = "ok"
            failure_stage = None
            failure_code = None
            try:
                _wide_initial, _wide_mutated, recovery_code = pridict2_sheet2_wide_sequences(record)
                design, locus_id, _window = canonicalize_pridict2_sheet2(record)
                designs.setdefault(design.design_id, design)
                _record_target_locus(target_loci, locus_id, _window)
                if recovery_code is not None:
                    recoveries.write(
                        {
                            "source_record_id": record_id,
                            "source_study": "PRIDICT2",
                            "source_sheet": "2",
                            "source_excel_row": excel_row,
                            "recovery_code": recovery_code,
                            "evidence_fields": (
                                "wide_initial_train|wide_mutated_train|spacer|RTT|"
                                "RT_initial_location|RT_mutated_location"
                            ),
                        }
                    )
                    counts[recovery_code] += 1
            except (KeyError, TypeError, ValueError) as exc:
                status = "quarantined"
                failure_stage = "canonicalization"
                failure_code = str(exc).split(":", 1)[0]
                _quarantine(
                    quarantine,
                    "PRIDICT2",
                    "2",
                    excel_row,
                    record_id,
                    "canonicalization",
                    exc,
                )
            counts["sheet2_designs"] += 1
            for outcome_column, condition in conditions2.items():
                if record.get(outcome_column) is None:
                    continue
                observation, label_error = _observation_record(
                    record_id=record_id,
                    design_id=design.design_id if design else None,
                    locus_id=locus_id,
                    condition=condition,
                    source_study="PRIDICT2",
                    source_dataset_id="PRIDICT2_Library_Diverse",
                    source_sheet="2",
                    source_excel_row=excel_row,
                    outcome_column=outcome_column,
                    outcome_value=record[outcome_column],
                    upstream_fold=record.get("testset_fold"),
                    normalization_status=status,
                    failure_stage=failure_stage,
                    failure_code=failure_code,
                )
                on_target.write(observation)
                counts["sheet2_observations"] += 1
                if label_error:
                    _quarantine(
                        quarantine,
                        "PRIDICT2",
                        "2",
                        excel_row,
                        record_id,
                        "label",
                        label_error,
                    )

        conditions7 = pridict2_sheet7_conditions()
        _headers, records = iter_sheet_records(workbook["7"], SheetSpec(1, 2), workbook_hash)
        for excel_row, record_id, record, _cells in records:
            design = None
            locus_id = None
            status = "ok"
            failure_stage = None
            failure_code = None
            try:
                design, locus_id, _window = canonicalize_pridict2_sheet7(record)
                designs.setdefault(design.design_id, design)
                _record_target_locus(target_loci, locus_id, _window)
            except (KeyError, TypeError, ValueError) as exc:
                status = "quarantined"
                failure_stage = "canonicalization"
                failure_code = str(exc).split(":", 1)[0]
                _quarantine(
                    quarantine,
                    "PRIDICT2",
                    "7",
                    excel_row,
                    record_id,
                    "canonicalization",
                    exc,
                )
            counts["sheet7_designs"] += 1
            for outcome_column, condition in conditions7.items():
                if record.get(outcome_column) is None:
                    continue
                observation, label_error = _observation_record(
                    record_id=record_id,
                    design_id=design.design_id if design else None,
                    locus_id=locus_id,
                    condition=condition,
                    source_study="PRIDICT2",
                    source_dataset_id="PRIDICT2_Arrayed_PE",
                    source_sheet="7",
                    source_excel_row=excel_row,
                    outcome_column=outcome_column,
                    outcome_value=record[outcome_column],
                    normalization_status=status,
                    failure_stage=failure_stage,
                    failure_code=failure_code,
                )
                on_target.write(observation)
                counts["sheet7_observations"] += 1
                if label_error:
                    _quarantine(
                        quarantine,
                        "PRIDICT2",
                        "7",
                        excel_row,
                        record_id,
                        "label",
                        label_error,
                    )
    finally:
        workbook.close()
    return {"counts": dict(counts)}


def _write_chromatin_task(
    path: Path,
    workbook_hash: str,
    destination: Path,
    quarantine: ParquetBatchWriter,
) -> dict[str, int]:
    frame = pd.read_excel(path, sheet_name="12", header=0)
    frame = frame.dropna(axis=1, how="all")
    frame.insert(
        0,
        "source_record_id",
        [source_record_id(workbook_hash, "12", row) for row in range(2, len(frame) + 2)],
    )
    frame.insert(1, "source_study", "PRIDICT2")
    frame.insert(2, "source_sheet", "12")
    frame.insert(3, "source_excel_row", range(2, len(frame) + 2))
    frame.insert(
        4,
        "observation_id",
        [
            observation_id(record_id, "PE_editing_efficiency")
            for record_id in frame["source_record_id"]
        ],
    )
    frame.insert(5, "outcome_type", "absolute_editing_efficiency")
    frame.insert(6, "outcome_source_column", "PE_editing_efficiency")
    frame.insert(7, "outcome_source_value", frame["PE_editing_efficiency"])
    frame.insert(8, "outcome_source_unit", "percent")
    labels = pd.to_numeric(frame["PE_editing_efficiency"], errors="coerce")
    invalid = labels.notna() & ~labels.between(0, 100, inclusive="both")
    frame.insert(9, "editing_efficiency_fraction", labels.mask(invalid) / 100.0)
    registration = label_registration("PRIDICT2", "12", "PE_editing_efficiency")
    for offset, field in enumerate(
        (
            "label_ontology_version",
            "label_definition_id",
            "label_compatibility_group",
            "label_evidence_status",
        ),
        start=10,
    ):
        frame.insert(offset, field, registration[field])
    frame.insert(
        14,
        "normalization_status",
        ["quarantined" if value else "ok" for value in invalid],
    )
    frame.insert(
        15,
        "failure_stage",
        ["label" if value else None for value in invalid],
    )
    frame.insert(
        16,
        "failure_code",
        ["absolute_efficiency_out_of_range" if value else None for value in invalid],
    )
    for index in frame.index[invalid]:
        _quarantine(
            quarantine,
            "PRIDICT2",
            "12",
            int(frame.at[index, "source_excel_row"]),
            str(frame.at[index, "source_record_id"]),
            "label",
            f"absolute_efficiency_out_of_range:{labels.at[index]}",
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    frame.to_parquet(destination, compression="zstd", index=False)
    return {
        "rows": len(frame),
        "pe_labels": int(frame["PE_editing_efficiency"].notna().sum()),
        "invalid_pe_labels": int(invalid.sum()),
        "columns": len(frame.columns),
    }


def _feature_schema(columns: list[str], *, context: bool = False) -> pa.Schema:
    integer_prefixes = (
        "edit_position",
        "edit_type_",
        "gc_count_",
        "max_run_",
        "poly_t_max_run",
        "poly_t_risk_",
        "pam_length",
        "pam_disrupted",
    )
    string_columns = {
        "pam_before",
        "pam_after",
        "base_before_edit",
        "base_after_edit",
        "base_before_nick",
        "base_after_nick",
        "cell_line",
        "editor_system",
        "pam_family",
        "scaffold_class",
        "organism",
        "mmr_status",
        "assay_type",
        "target_context",
        "delivery_context",
    }
    fields: list[tuple[str, pa.DataType]] = [
        (("observation_id" if context else "design_id"), pa.string())
    ]
    if context:
        fields.append(("design_id", pa.string()))
    for column in columns:
        if column in string_columns:
            dtype = pa.string()
        elif (
            column.endswith("_length")
            or column.startswith(integer_prefixes)
            or column == "right_homology_arm_length"
        ):
            dtype = pa.int64()
        else:
            dtype = pa.float64()
        fields.append((column, dtype))
    fields.extend([("feature_status", pa.string()), ("feature_failure_codes", pa.string())])
    return pa.schema(fields)


def _write_features(
    staging: Path,
    designs: dict[str, CanonicalDesign],
    on_target_path: Path,
    registry: dict[str, Any],
) -> dict[str, Any]:
    f0_columns = feature_columns(registry, "F0_core_v1")
    f1_columns = feature_columns(registry, "F1_sequence_v1")
    context_columns = feature_columns(registry, "F1_context_v1")
    f0_writer = ParquetBatchWriter(
        staging / "features" / "f0_core_v1.parquet",
        _feature_schema(f0_columns),
        batch_size=5_000,
    )
    f1_writer = ParquetBatchWriter(
        staging / "features" / "f1_sequence_v1.parquet",
        _feature_schema(f1_columns),
        batch_size=2_000,
    )
    feature_failure_counts = Counter()
    for design_id in sorted(designs):
        design = designs[design_id]
        f0 = compute_f0(design)
        f1, failures = compute_f1_sequence(design)
        for failure in failures:
            feature_failure_counts[failure] += 1
        status = "complete" if not failures else "complete_with_null_features"
        failure_text = "|".join(sorted(failures)) if failures else None
        f0_writer.write(
            {
                "design_id": design_id,
                **{column: f0[column] for column in f0_columns},
                "feature_status": "complete",
                "feature_failure_codes": None,
            }
        )
        f1_writer.write(
            {
                "design_id": design_id,
                **{column: f1[column] for column in f1_columns},
                "feature_status": status,
                "feature_failure_codes": failure_text,
            }
        )
    f0_writer.close()
    f1_writer.close()

    f1_frame = pd.read_parquet(staging / "features" / "f1_sequence_v1.parquet")
    observation_columns = [
        "observation_id",
        "design_id",
        *registry["feature_sets"]["F1_context_v1"]["features"],
    ]
    context_writer = ParquetBatchWriter(
        staging / "features" / "f1_context_v1.parquet",
        _feature_schema(context_columns, context=True),
        batch_size=2_000,
    )
    parquet_file = pq.ParquetFile(on_target_path)
    for batch in parquet_file.iter_batches(batch_size=20_000, columns=observation_columns):
        observations = batch.to_pandas()
        merged = observations.merge(f1_frame, on="design_id", how="left", validate="many_to_one")
        for record in merged.to_dict(orient="records"):
            failure_text = record.get("feature_failure_codes")
            status = record.get("feature_status")
            if pd.isna(record.get("design_id")):
                status = "quarantined"
                failure_text = "design_unavailable"
            output: dict[str, Any] = {
                "observation_id": record["observation_id"],
                "design_id": None if pd.isna(record.get("design_id")) else record["design_id"],
            }
            for column in context_columns:
                value = record.get(column)
                output[column] = None if pd.isna(value) else value
            output["feature_status"] = status
            output["feature_failure_codes"] = None if pd.isna(failure_text) else failure_text
            context_writer.write(output)
    context_writer.close()
    return {
        "unique_designs": len(designs),
        "f0_width": len(f0_columns),
        "f1_sequence_width": len(f1_columns),
        "f1_context_width": len(context_columns),
        "feature_failure_counts": dict(feature_failure_counts),
    }


def _write_designs(path: Path, designs: dict[str, CanonicalDesign]) -> int:
    writer = ParquetBatchWriter(path, DESIGN_SCHEMA, batch_size=5_000)
    for design_id in sorted(designs):
        writer.write(designs[design_id].as_record())
    writer.close()
    return writer.rows


def _write_target_loci(path: Path, target_loci: dict[str, str]) -> int:
    writer = ParquetBatchWriter(path, TARGET_LOCUS_SCHEMA, batch_size=5_000)
    for locus_id in sorted(target_loci):
        writer.write(
            {
                "target_locus_id": locus_id,
                "oriented_74nt_window": target_loci[locus_id],
            }
        )
    writer.close()
    return writer.rows


def _update_difference_metric(
    metrics: dict[str, dict[str, Any]],
    name: str,
    source: pd.Series,
    recalculated: pd.Series,
    *,
    source_scale: float = 1.0,
    tolerance: float = 1e-8,
) -> None:
    source_numeric = pd.to_numeric(source, errors="coerce") * source_scale
    recalculated_numeric = pd.to_numeric(recalculated, errors="coerce")
    valid = source_numeric.notna() & recalculated_numeric.notna()
    difference = (source_numeric[valid] - recalculated_numeric[valid]).abs()
    metric = metrics.setdefault(
        name,
        {
            "comparison_type": "numeric",
            "source_non_null": 0,
            "recalculated_non_null": 0,
            "compared": 0,
            "within_tolerance": 0,
            "sum_absolute_difference": 0.0,
            "max_absolute_difference": 0.0,
            "tolerance": tolerance,
        },
    )
    metric["source_non_null"] += int(source_numeric.notna().sum())
    metric["recalculated_non_null"] += int(recalculated_numeric.notna().sum())
    metric["compared"] += int(valid.sum())
    metric["within_tolerance"] += int(difference.le(tolerance).sum())
    metric["sum_absolute_difference"] += float(difference.sum())
    if not difference.empty:
        metric["max_absolute_difference"] = max(
            float(metric["max_absolute_difference"]), float(difference.max())
        )


def _update_categorical_metric(
    metrics: dict[str, dict[str, Any]],
    name: str,
    source: pd.Series,
    recalculated: pd.Series,
) -> None:
    source_values = source.astype("string").str.strip().str.upper()
    recalculated_values = recalculated.astype("string").str.strip().str.upper()
    valid = source_values.notna() & recalculated_values.notna()
    metric = metrics.setdefault(
        name,
        {
            "comparison_type": "categorical",
            "source_non_null": 0,
            "recalculated_non_null": 0,
            "compared": 0,
            "exact_match": 0,
        },
    )
    metric["source_non_null"] += int(source_values.notna().sum())
    metric["recalculated_non_null"] += int(recalculated_values.notna().sum())
    metric["compared"] += int(valid.sum())
    metric["exact_match"] += int(source_values[valid].eq(recalculated_values[valid]).sum())


def _write_parquet_frame(
    frame: pd.DataFrame,
    destination: Path,
    *,
    schema: pa.Schema | None = None,
    metadata: dict[str, str] | None = None,
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    table = pa.Table.from_pandas(frame, schema=schema, preserve_index=False)
    encoded = dict(table.schema.metadata or {})
    encoded.update(
        {str(key).encode(): str(value).encode() for key, value in (metadata or {}).items()}
    )
    table = table.replace_schema_metadata(encoded)
    pq.write_table(
        table,
        destination,
        compression="zstd",
        use_dictionary=True,
        write_statistics=True,
    )


def _write_source_label_audit_views(
    staging: Path,
    contract: dict[str, Any],
) -> dict[str, Any]:
    """Write source-label audit views; none is a unified-training entrypoint."""

    ontology = load_label_ontology()
    on_target = pd.read_parquet(staging / "tasks" / "pegRNA_on_target.parquet")
    on_target_eligible = (
        on_target["normalization_status"].eq("ok")
        & on_target["editing_efficiency_fraction"].notna()
        & on_target["design_id"].notna()
        & on_target["target_locus_id"].notna()
    )
    clean = on_target.loc[on_target_eligible].reset_index(drop=True)
    audit_view = contract["non_training_audit_views"]["pegRNA_on_target_clean_mixed_labels"]
    expected_clean = int(audit_view["expected_rows"])
    if len(clean) != expected_clean:
        raise ValueError(f"clean_mixed_label_view_count:{len(clean)}:{expected_clean}")
    _write_parquet_frame(
        clean,
        staging / str(audit_view["output"]),
        schema=ON_TARGET_SCHEMA,
        metadata={
            "training_authorized": "false",
            "label_ontology_version": ontology["ontology_version"],
            "warning": "mixed_incompatible_label_definitions",
        },
    )

    on_target_contract = contract["source_label_group_audit_views"]["pegRNA_on_target"]
    on_target_groups = on_target_contract["groups"]
    observed_groups = set(clean["label_compatibility_group"])
    expected_groups = set(on_target_groups)
    if observed_groups != expected_groups:
        raise ValueError(
            "on_target_label_groups_mismatch:"
            f"observed={sorted(observed_groups)}:expected={sorted(expected_groups)}"
        )
    group_entries: list[dict[str, Any]] = []
    on_target_counts: dict[str, int] = {}
    on_target_root = staging / str(on_target_contract["output_directory"])
    for group, definition in on_target_groups.items():
        group_frame = clean.loc[clean["label_compatibility_group"].eq(group)].reset_index(drop=True)
        expected_rows = int(definition["expected_rows"])
        if len(group_frame) != expected_rows:
            raise ValueError(f"label_group_count:{group}:{len(group_frame)}:{expected_rows}")
        definition_ids = sorted(group_frame["label_definition_id"].unique())
        destination = on_target_root / f"{group}.parquet"
        _write_parquet_frame(
            group_frame,
            destination,
            schema=ON_TARGET_SCHEMA,
            metadata={
                "training_authorized": "false",
                "purpose": "source_label_audit",
                "task": "pegRNA_on_target",
                "label_ontology_version": ontology["ontology_version"],
                "label_compatibility_group": group,
            },
        )
        on_target_counts[group] = len(group_frame)
        group_entries.append(
            {
                "task": "pegRNA_on_target",
                "label_compatibility_group": group,
                "label_definition_ids": definition_ids,
                "rows": len(group_frame),
                "path": str(destination.relative_to(staging)),
                "training_authorized": False,
            }
        )

    chromatin = pd.read_parquet(staging / "tasks" / "chromatin_context.parquet")
    chromatin_eligible = (
        chromatin["normalization_status"].eq("ok")
        & chromatin["editing_efficiency_fraction"].notna()
    )
    chromatin_clean = chromatin.loc[chromatin_eligible].reset_index(drop=True)
    chromatin_contract = contract["source_label_group_audit_views"]["chromatin_context"]
    chromatin_groups = chromatin_contract["groups"]
    observed_chromatin_groups = set(chromatin_clean["label_compatibility_group"])
    if observed_chromatin_groups != set(chromatin_groups):
        raise ValueError(
            "chromatin_label_groups_mismatch:"
            f"observed={sorted(observed_chromatin_groups)}:"
            f"expected={sorted(chromatin_groups)}"
        )
    chromatin_counts: dict[str, int] = {}
    chromatin_root = staging / str(chromatin_contract["output_directory"])
    for group, definition in chromatin_groups.items():
        group_frame = chromatin_clean.loc[
            chromatin_clean["label_compatibility_group"].eq(group)
        ].reset_index(drop=True)
        expected_rows = int(definition["expected_rows"])
        if len(group_frame) != expected_rows:
            raise ValueError(f"label_group_count:{group}:{len(group_frame)}:{expected_rows}")
        destination = chromatin_root / f"{group}.parquet"
        _write_parquet_frame(
            group_frame,
            destination,
            metadata={
                "training_authorized": "false",
                "purpose": "source_label_audit",
                "task": "chromatin_context",
                "label_ontology_version": ontology["ontology_version"],
                "label_compatibility_group": group,
            },
        )
        chromatin_counts[group] = len(group_frame)
        group_entries.append(
            {
                "task": "chromatin_context",
                "label_compatibility_group": group,
                "label_definition_ids": sorted(group_frame["label_definition_id"].unique()),
                "rows": len(group_frame),
                "path": str(destination.relative_to(staging)),
                "training_authorized": False,
            }
        )

    index = {
        "format": "prime_editing_source_label_group_audit_index.v1",
        "label_ontology_version": ontology["ontology_version"],
        "training_authorized": False,
        "unified_training_label_required": "harmonized_editing_efficiency_fraction",
        "entries": group_entries,
    }
    index_path = staging / "reports" / "source_label_groups" / "index.json"
    index_path.parent.mkdir(parents=True, exist_ok=True)
    index_path.write_text(
        json.dumps(index, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    source_tuples = pd.concat(
        [
            on_target[
                [
                    "source_study",
                    "source_sheet",
                    "outcome_source_column",
                    "label_definition_id",
                    "label_compatibility_group",
                    "label_evidence_status",
                ]
            ],
            chromatin[
                [
                    "source_study",
                    "source_sheet",
                    "outcome_source_column",
                    "label_definition_id",
                    "label_compatibility_group",
                    "label_evidence_status",
                ]
            ],
        ],
        ignore_index=True,
    ).drop_duplicates()
    source_tuple_records = source_tuples.sort_values(
        ["source_study", "source_sheet", "outcome_source_column"]
    ).to_dict(orient="records")
    if len(source_tuple_records) != int(ontology["expected_source_mapping_count"]):
        raise ValueError(
            "resolved_label_mapping_count:"
            f"{len(source_tuple_records)}:{ontology['expected_source_mapping_count']}"
        )
    audit = {
        "status": "pass",
        "label_ontology_version": ontology["ontology_version"],
        "numeric_scale_semantics": "percent_divided_by_100_only",
        "source_label_training_authorized": False,
        "unified_training_resolution": "raw_reads_v1",
        "resolved_source_mapping_count": len(source_tuple_records),
        "resolved_source_mappings": source_tuple_records,
        "on_target": {
            "source_rows": len(on_target),
            "clean_mixed_audit_rows": len(clean),
            "excluded_rows": int((~on_target_eligible).sum()),
            "source_group_audit_rows": on_target_counts,
        },
        "chromatin_context": {
            "source_rows": len(chromatin),
            "excluded_missing_or_invalid_label_rows": int((~chromatin_eligible).sum()),
            "source_group_audit_rows": chromatin_counts,
        },
    }
    reports_root = staging / "reports"
    reports_root.mkdir(parents=True, exist_ok=True)
    (reports_root / "label_ontology_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return audit


def _write_source_recalculation_report(
    staging: Path,
    destination: Path,
) -> dict[str, Any]:
    observations = pd.read_parquet(
        staging / "tasks" / "pegRNA_on_target.parquet",
        columns=["source_record_id", "source_study", "source_sheet", "design_id"],
    )
    source_design = (
        observations.dropna(subset=["design_id"])
        .drop_duplicates(subset=["source_record_id"])
        .reset_index(drop=True)
    )
    f0 = pd.read_parquet(staging / "features" / "f0_core_v1.parquet").drop(
        columns=["feature_status", "feature_failure_codes"]
    )
    f1_columns = [
        "design_id",
        "tm_pbs_rna_dna",
        "tm_rtt_dna_dna_original",
        "tm_rtt_dna_dna_mismatch",
        "tm_rtt_rna_dna",
        "mfe_extension_plain",
        "mfe_extension_deepprime_poly_u6",
        "mfe_spacer_canonical",
        "mfe_spacer_u6",
        "mfe_rtt",
        "mfe_pbs",
    ]
    features = f0.merge(
        pd.read_parquet(
            staging / "features" / "f1_sequence_v1.parquet",
            columns=f1_columns,
        ),
        on="design_id",
        validate="one_to_one",
    )
    features = features.merge(
        pd.read_parquet(
            staging / "tables" / "designs.parquet",
            columns=[
                "design_id",
                "pam_before",
                "pam_after",
                "base_before_edit",
                "base_after_edit",
                "base_before_nick",
                "base_after_nick",
            ],
        ),
        on="design_id",
        validate="one_to_one",
    )
    metrics: dict[str, dict[str, Any]] = {}
    mappings: dict[tuple[str, str], dict[str, tuple[str, float, float]]] = {
        ("DeepPrime", "all"): {
            "PBSlen": ("pbs_length", 1.0, 0.0),
            "RTlen": ("rtt_length", 1.0, 0.0),
            "RT-PBSlen": ("extension_length", 1.0, 0.0),
            "Edit_pos": ("edit_position", 1.0, 0.0),
            "Edit_len": ("edit_length", 1.0, 0.0),
            "RHA_len": ("right_homology_arm_length", 1.0, 0.0),
            "type_sub": ("edit_type_substitution", 1.0, 0.0),
            "type_ins": ("edit_type_insertion", 1.0, 0.0),
            "type_del": ("edit_type_deletion", 1.0, 0.0),
            "nGCcnt1": ("gc_count_pbs", 1.0, 0.0),
            "nGCcnt2": ("gc_count_rtt", 1.0, 0.0),
            "nGCcnt3": ("gc_count_extension", 1.0, 0.0),
            "fGCcont1": ("gc_fraction_pbs", 0.01, 1e-10),
            "fGCcont2": ("gc_fraction_rtt", 0.01, 1e-10),
            "fGCcont3": ("gc_fraction_extension", 0.01, 1e-10),
            "Tm1": ("tm_pbs_rna_dna", 1.0, 1e-8),
            "Tm2new": ("tm_rtt_dna_dna_original", 1.0, 1e-8),
            "Tm3": ("tm_rtt_dna_dna_mismatch", 1.0, 1e-8),
            "Tm4": ("tm_rtt_rna_dna", 1.0, 1e-8),
            "MFE3": ("mfe_extension_deepprime_poly_u6", 1.0, 0.051),
            "MFE4": ("mfe_spacer_u6", 1.0, 0.051),
        },
        ("PRIDICT2", "2"): {
            "PBSlength": ("pbs_length", 1.0, 0.0),
            "RTTlength": ("rtt_length", 1.0, 0.0),
            "Editing_Position": ("edit_position", 1.0, 0.0),
            "Correction_Length": ("edit_length", 1.0, 0.0),
            "RTToverhanglength": ("right_homology_arm_length", 1.0, 0.0),
            "Spacer_GC_count": ("gc_count_spacer", 1.0, 0.0),
            "PBS_GC_count": ("gc_count_pbs", 1.0, 0.0),
            "RT_GC_count": ("gc_count_rtt", 1.0, 0.0),
            "Extension_GC_count": ("gc_count_extension", 1.0, 0.0),
            "Spacer_GC_content": ("gc_fraction_spacer", 0.01, 1e-10),
            "PBS_GC_content": ("gc_fraction_pbs", 0.01, 1e-10),
            "RT_GC_content": ("gc_fraction_rtt", 0.01, 1e-10),
            "Extension_GC_content": ("gc_fraction_extension", 0.01, 1e-10),
            "PBSmt": ("tm_pbs_rna_dna", 1.0, 1e-8),
            "RTmt": ("tm_rtt_rna_dna", 1.0, 1e-8),
            "MFE_extension": ("mfe_extension_plain", 1.0, 0.051),
            "MFE_protospacer": ("mfe_spacer_canonical", 1.0, 0.051),
            "MFE_rt": ("mfe_rtt", 1.0, 0.051),
            "MFE_pbs": ("mfe_pbs", 1.0, 0.051),
        },
        ("PRIDICT2", "7"): {
            "PBSlength": ("pbs_length", 1.0, 0.0),
            "RTTlength": ("rtt_length", 1.0, 0.0),
            "Edit_position": ("edit_position", 1.0, 0.0),
            "Edit_length": ("edit_length", 1.0, 0.0),
            "RTToverhanglength": ("right_homology_arm_length", 1.0, 0.0),
        },
    }

    for (source, sheet_key), field_mappings in mappings.items():
        sheets = [str(index) for index in range(1, 21)] if source == "DeepPrime" else [sheet_key]
        for sheet in sheets:
            raw_path = staging / "raw_mirror" / f"{source.lower()}__{sheet}.parquet"
            available = set(pq.read_schema(raw_path).names)
            numeric_source_columns = [column for column in field_mappings if column in available]
            categorical_mappings = (
                {
                    "PAMbefore": "pam_before",
                    "PAMafter": "pam_after",
                    "basebefore": "base_before_edit",
                    "baseafter": "base_after_edit",
                    "nickbasebefore": "base_before_nick",
                    "nickbaseafter": "base_after_nick",
                }
                if source == "PRIDICT2" and sheet == "2"
                else {}
            )
            categorical_source_columns = [
                column for column in categorical_mappings if column in available
            ]
            raw = pd.read_parquet(
                raw_path,
                columns=[
                    "__source_record_id",
                    *numeric_source_columns,
                    *categorical_source_columns,
                ],
            ).rename(columns={"__source_record_id": "source_record_id"})
            mapping = source_design[
                source_design["source_study"].eq(source) & source_design["source_sheet"].eq(sheet)
            ][["source_record_id", "design_id"]]
            joined = raw.merge(mapping, on="source_record_id", validate="one_to_one").merge(
                features, on="design_id", validate="many_to_one"
            )
            for source_column in numeric_source_columns:
                feature, scale, tolerance = field_mappings[source_column]
                source_values = joined[source_column]
                metric_source_column = source_column
                if source == "PRIDICT2" and sheet == "2" and source_column == ("Editing_Position"):
                    source_values = pd.to_numeric(source_values, errors="coerce") + 1
                    metric_source_column = "Editing_Position_plus_1"
                _update_difference_metric(
                    metrics,
                    f"{source}.{metric_source_column}__vs__{feature}",
                    source_values,
                    joined[feature],
                    source_scale=scale,
                    tolerance=tolerance,
                )
            for source_column in categorical_source_columns:
                feature = categorical_mappings[source_column]
                recalculated_values = joined[feature]
                if source_column in {"PAMbefore", "PAMafter"}:
                    recalculated_values = recalculated_values.astype("string").str[-2:]
                _update_categorical_metric(
                    metrics,
                    f"{source}.{source_column}__vs__{feature}",
                    joined[source_column],
                    recalculated_values,
                )

    for metric in metrics.values():
        if metric["comparison_type"] == "categorical":
            metric["mismatch_count"] = int(metric["compared"]) - int(metric["exact_match"])
            continue
        compared = int(metric["compared"])
        metric["mean_absolute_difference"] = (
            float(metric["sum_absolute_difference"]) / compared if compared else None
        )
        metric["mismatch_count"] = compared - int(metric["within_tolerance"])
        del metric["sum_absolute_difference"]
    report = {
        "policy": "source values are audit-only and never copied into F0/F1",
        "interpretation": {
            "length_gc_geometry": "expected to agree after direction/unit normalization",
            "deepprime_tm": "source defaults and legacy Tm3 behavior differ from frozen v1 parameters",
            "pridict2_tm": "source values use Wallace Tm and are not definition-compatible with v1 Tm_NN",
            "pridict2_spacer": "source spacer may carry a U6-leading-G adjustment; v1 restores the target-matched spacer",
            "pridict2_base_after_edit": (
                "for multi-base replacements, source baseafter is adjacent to the "
                "first edited base, while unified base_after_edit is the base after "
                "the complete replaced span; this registered semantic difference is "
                "not a sequence reconstruction failure"
            ),
            "mfe": "source values are rounded and may use another ViennaRNA version",
        },
        "metrics": metrics,
    }
    destination.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def _hamming_one_target_pairs(
    deepprime_windows: dict[str, str],
    pridict2_windows: dict[str, str],
) -> list[dict[str, Any]]:
    """Find deterministic cross-source Hamming-distance-1 target pairs."""

    deepprime_by_window = {window: locus_id for locus_id, window in deepprime_windows.items()}
    pairs: list[dict[str, Any]] = []
    for pridict_locus_id in sorted(pridict2_windows):
        pridict_window = pridict2_windows[pridict_locus_id]
        for index, pridict_base in enumerate(pridict_window):
            for deepprime_base in "ACGT":
                if deepprime_base == pridict_base:
                    continue
                candidate = pridict_window[:index] + deepprime_base + pridict_window[index + 1 :]
                deepprime_locus_id = deepprime_by_window.get(candidate)
                if deepprime_locus_id is None:
                    continue
                pairs.append(
                    {
                        "deepprime_target_locus_id": deepprime_locus_id,
                        "pridict2_target_locus_id": pridict_locus_id,
                        "hamming_distance": 1,
                        "differing_position_1based": index + 1,
                        "deepprime_base": deepprime_base,
                        "pridict2_base": pridict_base,
                    }
                )
    return sorted(
        pairs,
        key=lambda row: (
            row["deepprime_target_locus_id"],
            row["pridict2_target_locus_id"],
        ),
    )


def _write_overlap_report(
    on_target_path: Path,
    target_loci_path: Path,
    destination: Path,
) -> dict[str, Any]:
    frame = pd.read_parquet(
        on_target_path,
        columns=[
            "observation_id",
            "design_id",
            "target_locus_id",
            "condition_id",
            "source_record_id",
            "source_study",
            "source_sheet",
            "source_excel_row",
            "outcome_type",
        ],
    )
    valid = frame.dropna(subset=["design_id"])
    designs_by_source = {
        source: set(group["design_id"]) for source, group in valid.groupby("source_study")
    }
    targets_by_source = {
        source: set(group["target_locus_id"].dropna())
        for source, group in valid.groupby("source_study")
    }
    deepprime_designs = designs_by_source.get("DeepPrime", set())
    pridict_designs = designs_by_source.get("PRIDICT2", set())
    deepprime_targets = targets_by_source.get("DeepPrime", set())
    pridict_targets = targets_by_source.get("PRIDICT2", set())
    exact_designs = deepprime_designs & pridict_designs
    exact_targets = deepprime_targets & pridict_targets
    report_columns = list(frame.columns)
    frame[frame["design_id"].isin(exact_designs)][report_columns].to_parquet(
        destination.parent / "cross_source_exact_design_observations.parquet",
        compression="zstd",
        index=False,
    )
    frame[frame["target_locus_id"].isin(exact_targets)][report_columns].to_parquet(
        destination.parent / "cross_source_exact_target_observations.parquet",
        compression="zstd",
        index=False,
    )
    frame[frame["observation_id"].duplicated(keep=False)][report_columns].to_parquet(
        destination.parent / "duplicate_observation_ids.parquet",
        compression="zstd",
        index=False,
    )
    target_loci = pd.read_parquet(target_loci_path).set_index("target_locus_id")[
        "oriented_74nt_window"
    ]
    deepprime_windows = {locus_id: target_loci[locus_id] for locus_id in sorted(deepprime_targets)}
    pridict2_windows = {locus_id: target_loci[locus_id] for locus_id in sorted(pridict_targets)}
    near_pairs = _hamming_one_target_pairs(deepprime_windows, pridict2_windows)
    near_frame = pd.DataFrame(
        near_pairs,
        columns=[
            "deepprime_target_locus_id",
            "pridict2_target_locus_id",
            "hamming_distance",
            "differing_position_1based",
            "deepprime_base",
            "pridict2_base",
        ],
    )
    near_frame.to_parquet(
        destination.parent / "cross_source_near_target_pairs.parquet",
        compression="zstd",
        index=False,
    )
    report = {
        "observations": len(frame),
        "observations_with_design": len(valid),
        "exact_duplicate_observation_ids": int(frame["observation_id"].duplicated().sum()),
        "unique_designs_by_source": {
            source: len(values) for source, values in designs_by_source.items()
        },
        "unique_targets_by_source": {
            source: len(values) for source, values in targets_by_source.items()
        },
        "cross_source_exact_design_overlap": len(exact_designs),
        "cross_source_exact_target_overlap": len(exact_targets),
        "row_level_reports": {
            "exact_design": "cross_source_exact_design_observations.parquet",
            "exact_target": "cross_source_exact_target_observations.parquet",
            "near_target": "cross_source_near_target_pairs.parquet",
            "duplicate_observation_id": "duplicate_observation_ids.parquet",
        },
        "near_duplicate_report_status": "computed",
        "near_duplicate_definition": (
            "cross-source canonical oriented 74-nt target windows with equal length "
            "and Hamming distance exactly 1"
        ),
        "cross_source_near_target_pairs": len(near_pairs),
        "deduplication_action": "none_observations_retained",
    }
    destination.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def _file_hashes(root: Path) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.name != "manifest.json":
            relative = str(path.relative_to(root))
            result[relative] = {"bytes": path.stat().st_size, "sha256": sha256(path)}
    return result


def _package_readme(version: str) -> str:
    return f"""# Prime editing unified data package {version}

This directory is generated from the verified DeepPrime Table S2 and PRIDICT2
supplementary workbook. The source workbooks are immutable and are not copied
into this package.

- `raw_mirror/`: source values by workbook and sheet; values are represented as
  strings while canonical task tables retain typed numerical columns.
- `tables/`: canonical designs and source lineage.
- `tasks/`: semantically separate on-target, mismatch/off-target, chromatin, and
  supporting-metadata views.
- `features/`: F0 core, sequence-only F1, and context-enriched F1.
- `quarantine/`: row-level failures retained with stable source identifiers.
- `reports/`: raw, normalization, feature, and overlap audits.
- `data_contract_v2.yaml`, `data_dictionary_v2.yaml`,
  `label_ontology_v1.yaml`, `label_harmonization_v1.yaml`, and
  `feature_registry_v1.yaml`: frozen schemas, label meanings, feature lists,
  algorithms, parameters, and exclusions.

Source identifiers, labels, upstream folds/models, and prediction scores are
lineage-only and are never registered as trainable features.

`editing_efficiency_fraction` standardizes only the numerical percent-to-
fraction scale. It remains a source-label audit field. Source-definition groups
under `reports/source_label_groups/` are also audit-only and have
`training_authorized=false`; they are not the final modelling interface.

Unified DeepPrime + PRIDICT2 training requires the separately named
`harmonized_editing_efficiency_fraction`, recomputed from raw reads according to
`label_harmonization_v1.yaml`. The workbook-only build records
`raw_recomputation_required` and intentionally does not materialize a unified
training table. Once raw counting passes all checks, the training view is
`tasks/pegRNA_on_target_harmonized_training.parquet`. The complete
source-preserving task remains in `tasks/pegRNA_on_target.parquet`.

DeepPrime mismatch sheets contain no sequence fields. Consequently,
`design_id` and `target_locus_id` are explicitly null in
`tasks/mismatch_off_target.parquet`; they are not guessed. Use the deterministic
`source_target_group_id` only for leakage grouping within that task. It is not a
canonical sequence identifier and cannot join F0/F1.

Patch releases may use exact redundant source fields to recover missing target
windows. Every such recovery is recorded in
`reports/deterministic_recoveries.parquet`. Approximate target matches are
report-only Hamming-distance-1 candidates and are never merged automatically.
Immutable source anomalies are count-locked in the contract and verified in
`reports/source_anomalies.json`; original workbooks and raw rows are never
silently repaired or deduplicated.

The package is verified with:

```text
python -m primetab.features verify --package data/prime_editing_unified/{version}
```
"""


def build_package(
    deepprime_path: Path,
    pridict2_path: Path,
    output_root: Path,
    version: str = "v1",
) -> Path:
    contract = load_contract()
    registry = load_feature_registry()
    dictionary = _load_yaml(DATA_DICTIONARY_PATH)
    ontology = load_label_ontology()
    harmonization = load_harmonization_contract()
    expected_version = str(contract["package_version"])
    expected_package = f"{contract['package_name']}/{expected_version}"
    if version != expected_version:
        raise ValueError(
            f"package_version_mismatch:requested={version}:contract={expected_version}"
        )
    if dictionary["package"] != expected_package:
        raise ValueError(
            f"data_dictionary_package_mismatch:{dictionary['package']}:{expected_package}"
        )
    if ontology["ontology_version"] != contract["label_ontology"]["version"]:
        raise ValueError(
            "label_ontology_version_mismatch:"
            f"{ontology['ontology_version']}:{contract['label_ontology']['version']}"
        )
    if harmonization["harmonization_version"] != contract["harmonized_label"]["version"]:
        raise ValueError(
            "label_harmonization_version_mismatch:"
            f"{harmonization['harmonization_version']}:"
            f"{contract['harmonized_label']['version']}"
        )
    final = output_root / version
    if final.exists():
        raise FileExistsError(f"Refusing to overwrite existing package: {final}")
    output_root.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{version}.staging-", dir=output_root))
    try:
        raw_audit = audit_sources(deepprime_path, pridict2_path)
        deepprime_hash = raw_audit["sources"]["deepprime"]["sha256"]
        pridict2_hash = raw_audit["sources"]["pridict2"]["sha256"]
        summary = load_deepprime_summary(deepprime_path)
        lineage = ParquetBatchWriter(staging / "tables" / "lineage.parquet", LINEAGE_SCHEMA)
        deep_profiles, deep_raw_rows = _scan_source(
            "DeepPrime",
            deepprime_path,
            deepprime_hash,
            deepprime_specs(summary),
            staging / "raw_mirror",
            lineage,
        )
        pridict_profiles, pridict_raw_rows = _scan_source(
            "PRIDICT2",
            pridict2_path,
            pridict2_hash,
            pridict2_specs(pridict2_path),
            staging / "raw_mirror",
            lineage,
        )
        lineage.close()

        on_target = ParquetBatchWriter(
            staging / "tasks" / "pegRNA_on_target.parquet", ON_TARGET_SCHEMA
        )
        mismatch = ParquetBatchWriter(
            staging / "tasks" / "mismatch_off_target.parquet", MISMATCH_SCHEMA
        )
        quarantine = ParquetBatchWriter(staging / "quarantine" / "rows.parquet", QUARANTINE_SCHEMA)
        recoveries = ParquetBatchWriter(
            staging / "reports" / "deterministic_recoveries.parquet",
            RECOVERY_SCHEMA,
        )
        designs: dict[str, CanonicalDesign] = {}
        target_loci: dict[str, str] = {}
        deep_result = _process_deepprime(
            deepprime_path,
            deepprime_hash,
            summary,
            on_target,
            mismatch,
            quarantine,
            designs,
            target_loci,
        )
        pridict_result = _process_pridict2(
            pridict2_path,
            pridict2_hash,
            on_target,
            quarantine,
            designs,
            target_loci,
            recoveries,
        )
        on_target.close()
        mismatch.close()
        recoveries.close()

        chromatin = _write_chromatin_task(
            pridict2_path,
            pridict2_hash,
            staging / "tasks" / "chromatin_context.parquet",
            quarantine,
        )
        source_label_audit = _write_source_label_audit_views(staging, contract)
        quarantine.close()
        source_profiles = deep_profiles + pridict_profiles
        source_anomalies = _validate_known_source_anomalies(source_profiles, contract)
        catalog = pd.DataFrame(source_profiles)
        catalog[
            [
                "source_study",
                "sheet",
                "role",
                "columns",
                "rows",
                "formula_cells",
                "formula_cells_with_cached_value",
                "formula_cells_missing_cached_value",
                "error_cells",
                "exact_duplicate_rows",
                "exact_duplicate_groups",
            ]
        ].to_parquet(
            staging / "tasks" / "supporting_metadata_catalog.parquet",
            compression="zstd",
            index=False,
        )
        designs_count = _write_designs(staging / "tables" / "designs.parquet", designs)
        target_loci_count = _write_target_loci(
            staging / "tables" / "target_loci.parquet", target_loci
        )
        feature_audit = _write_features(
            staging,
            designs,
            staging / "tasks" / "pegRNA_on_target.parquet",
            registry,
        )
        designs.clear()
        gc.collect()

        reports_root = staging / "reports"
        reports_root.mkdir(parents=True, exist_ok=True)
        raw_report = {
            **raw_audit,
            "sheet_profiles": deep_profiles + pridict_profiles,
            "raw_mirror_rows": {
                "DeepPrime": deep_raw_rows,
                "PRIDICT2": pridict_raw_rows,
            },
        }
        (reports_root / "raw_audit.json").write_text(
            json.dumps(raw_report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        (reports_root / "source_anomalies.json").write_text(
            json.dumps(source_anomalies, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        quarantine_frame = pd.read_parquet(
            staging / "quarantine" / "rows.parquet",
            columns=["source_study", "failure_stage", "failure_code"],
        )
        quarantine_by_stage_code = [
            {
                "failure_stage": stage,
                "failure_code": code,
                "rows": int(rows),
            }
            for (stage, code), rows in quarantine_frame.groupby(
                ["failure_stage", "failure_code"], dropna=False
            )
            .size()
            .items()
        ]
        normalization_audit = {
            "deepprime": deep_result,
            "pridict2": pridict_result,
            "chromatin_context": chromatin,
            "task_rows": {
                "pegRNA_on_target": on_target.rows,
                "pegRNA_on_target_clean_mixed_labels": source_label_audit["on_target"][
                    "clean_mixed_audit_rows"
                ],
                "mismatch_off_target": mismatch.rows,
                "chromatin_context": chromatin["rows"],
            },
            "source_label_audit_views": source_label_audit,
            "harmonized_training_views": {
                "status": "raw_recomputation_required",
                "label": "harmonized_editing_efficiency_fraction",
            },
            "source_anomalies": source_anomalies,
            "quarantine_rows": quarantine.rows,
            "quarantine_by_source": {
                str(source): int(rows)
                for source, rows in quarantine_frame.groupby("source_study", dropna=False)
                .size()
                .items()
            },
            "quarantine_by_stage_code": quarantine_by_stage_code,
            "unique_designs": designs_count,
            "unique_target_loci": target_loci_count,
            "deterministic_recovery_rows": recoveries.rows,
            "repair_policy": "strict_redundant_source_fields_only_no_heuristics",
        }
        (reports_root / "normalization_audit.json").write_text(
            json.dumps(normalization_audit, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        (reports_root / "feature_audit.json").write_text(
            json.dumps(feature_audit, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        source_recalculation = _write_source_recalculation_report(
            staging,
            reports_root / "source_recalculation_audit.json",
        )
        overlap = _write_overlap_report(
            staging / "tasks" / "pegRNA_on_target.parquet",
            staging / "tables" / "target_loci.parquet",
            reports_root / "overlap_audit.json",
        )
        shutil.copy2(DATA_CONTRACT_PATH, staging / "data_contract_v2.yaml")
        shutil.copy2(FEATURE_REGISTRY_PATH, staging / "feature_registry_v1.yaml")
        shutil.copy2(DATA_DICTIONARY_PATH, staging / "data_dictionary_v2.yaml")
        shutil.copy2(LABEL_ONTOLOGY_PATH, staging / "label_ontology_v1.yaml")
        shutil.copy2(
            HARMONIZATION_CONTRACT_PATH,
            staging / "label_harmonization_v1.yaml",
        )
        (staging / "README.md").write_text(_package_readme(version), encoding="utf-8")

        manifest = {
            "package_name": contract["package_name"],
            "package_version": version,
            "contract_version": contract.get("contract_version", f"v{contract['version']}"),
            "dictionary_version": dictionary.get("dictionary_version", f"v{dictionary['version']}"),
            "label_ontology_version": load_label_ontology()["ontology_version"],
            "label_harmonization_version": harmonization["harmonization_version"],
            "harmonized_label_status": "raw_recomputation_required",
            "generated_at_utc": datetime.now(UTC).isoformat(),
            "inputs": raw_audit["sources"],
            "runtime": {
                "python": platform.python_version(),
                "platform": platform.platform(),
                "biopython": importlib.metadata.version("biopython"),
                "viennarna": importlib.metadata.version("ViennaRNA"),
                "pandas": pd.__version__,
                "pyarrow": pa.__version__,
                "openpyxl": importlib.metadata.version("openpyxl"),
            },
            "counts": normalization_audit,
            "feature_widths": feature_audit,
            "source_recalculation": source_recalculation,
            "overlap": overlap,
            "output_files": _file_hashes(staging),
            "manifest_identity_excludes": ["generated_at_utc"],
        }
        (staging / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

        verification = verify_package(staging, package_label=version)
        (reports_root / "verification.json").write_text(
            json.dumps(verification, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        # Refresh manifest hashes after the verification report is added.
        manifest["output_files"] = _file_hashes(staging)
        (staging / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        if final.exists():
            raise FileExistsError(f"Refusing to overwrite existing package: {final}")
        staging.replace(final)
        return final
    except Exception as exc:
        (staging / "BUILD_FAILED.txt").write_text(
            f"{exc.__class__.__name__}: {exc}\n", encoding="utf-8"
        )
        raise


def _check_feature_file(
    package: Path,
    relative: str,
    expected: list[str],
    id_columns: set[str],
) -> list[str]:
    schema = pq.read_schema(package / relative)
    actual = [name for name in schema.names if name not in id_columns]
    actual = [name for name in actual if name not in {"feature_status", "feature_failure_codes"}]
    if actual != expected:
        raise ValueError(
            f"feature_schema_mismatch:{relative}:expected={len(expected)}:actual={len(actual)}"
        )
    return actual


def verify_package(
    package: Path,
    *,
    package_label: str | None = None,
) -> dict[str, Any]:
    manifest_path = package / "manifest.json"
    if not manifest_path.exists():
        raise FileNotFoundError(f"Missing manifest: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    registry = _load_yaml(package / "feature_registry_v1.yaml")
    contract = _load_yaml(package / "data_contract_v2.yaml")
    dictionary = _load_yaml(package / "data_dictionary_v2.yaml")
    ontology = _load_yaml(package / "label_ontology_v1.yaml")
    harmonization = _load_yaml(package / "label_harmonization_v1.yaml")
    checks: list[dict[str, Any]] = []

    expected_package = f"{contract['package_name']}/{contract['package_version']}"
    metadata_checks = {
        "manifest_matches_contract_package_name": (
            manifest["package_name"] == contract["package_name"]
        ),
        "manifest_matches_contract_package_version": (
            manifest["package_version"] == contract["package_version"]
        ),
        "dictionary_matches_contract_package": dictionary["package"] == expected_package,
        "manifest_matches_label_ontology_version": (
            manifest["label_ontology_version"] == ontology["ontology_version"]
        ),
        "contract_matches_label_ontology_version": (
            contract["label_ontology"]["version"] == ontology["ontology_version"]
        ),
        "manifest_matches_label_harmonization_version": (
            manifest["label_harmonization_version"] == harmonization["harmonization_version"]
        ),
        "contract_matches_label_harmonization_version": (
            contract["harmonized_label"]["version"] == harmonization["harmonization_version"]
        ),
        "workbook_build_requires_raw_recomputation": (
            manifest["harmonized_label_status"] == "raw_recomputation_required"
        ),
    }
    for name, passed in metadata_checks.items():
        checks.append({"check": f"metadata:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"metadata_invariant_failed:{name}")

    for relative, expected in manifest["output_files"].items():
        path = package / relative
        passed = (
            path.is_file()
            and path.stat().st_size == expected["bytes"]
            and sha256(path) == expected["sha256"]
        )
        checks.append({"check": f"file_hash:{relative}", "passed": passed})
        if not passed:
            raise ValueError(f"output_file_hash_mismatch:{relative}")

    declared_files = set(manifest["output_files"])
    actual_files = {
        str(path.relative_to(package))
        for path in package.rglob("*")
        if path.is_file() and path.name != "manifest.json"
    }
    passed = actual_files == declared_files
    checks.append({"check": "file_inventory:exact", "passed": passed})
    if not passed:
        raise ValueError(
            "output_file_inventory_mismatch:"
            f"missing={sorted(declared_files - actual_files)}:"
            f"extra={sorted(actual_files - declared_files)}"
        )

    task_counts = {
        "pegRNA_on_target": pq.ParquetFile(
            package / "tasks" / "pegRNA_on_target.parquet"
        ).metadata.num_rows,
        "mismatch_off_target": pq.ParquetFile(
            package / "tasks" / "mismatch_off_target.parquet"
        ).metadata.num_rows,
        "chromatin_context": pq.ParquetFile(
            package / "tasks" / "chromatin_context.parquet"
        ).metadata.num_rows,
    }
    expected_counts = {
        name: int(
            definition["expected_observations"]
            if "expected_observations" in definition
            else definition["expected_rows"]
        )
        for name, definition in contract["task_views"].items()
    }
    for name, count in task_counts.items():
        passed = count == expected_counts[name]
        checks.append({"check": f"task_count:{name}", "passed": passed, "value": count})
        if not passed:
            raise ValueError(f"task_count_mismatch:{name}:{count}:{expected_counts[name]}")

    mixed_definition = contract["non_training_audit_views"]["pegRNA_on_target_clean_mixed_labels"]
    mixed_path = package / str(mixed_definition["output"])
    mixed_rows = pq.ParquetFile(mixed_path).metadata.num_rows
    expected_mixed_rows = int(mixed_definition["expected_rows"])
    mixed_metadata = pq.read_metadata(mixed_path).metadata or {}
    mixed_checks = {
        "expected_rows": mixed_rows == expected_mixed_rows,
        "training_not_authorized": mixed_metadata.get(b"training_authorized") == b"false",
        "mixed_warning_present": (
            mixed_metadata.get(b"warning") == b"mixed_incompatible_label_definitions"
        ),
    }
    for name, passed in mixed_checks.items():
        checks.append({"check": f"mixed_label_audit_view:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"mixed_label_audit_view_failed:{name}")

    source_group_paths: dict[str, Path] = {}
    source_group_counts: dict[str, int] = {}
    for task_name in ("pegRNA_on_target", "chromatin_context"):
        definition = contract["source_label_group_audit_views"][task_name]
        root = package / str(definition["output_directory"])
        for group, group_definition in definition["groups"].items():
            group_path = root / f"{group}.parquet"
            rows = pq.ParquetFile(group_path).metadata.num_rows
            expected_rows = int(group_definition["expected_rows"])
            metadata = pq.read_metadata(group_path).metadata or {}
            group_checks = {
                "expected_rows": rows == expected_rows,
                "training_not_authorized": metadata.get(b"training_authorized") == b"false",
                "source_audit_purpose": metadata.get(b"purpose") == b"source_label_audit",
                "task_metadata": metadata.get(b"task") == task_name.encode(),
                "group_metadata": metadata.get(b"label_compatibility_group") == group.encode(),
            }
            for name, passed in group_checks.items():
                checks.append(
                    {
                        "check": f"source_label_group_audit:{task_name}:{group}:{name}",
                        "passed": passed,
                    }
                )
                if not passed:
                    raise ValueError(f"source_label_group_audit_failed:{task_name}:{group}:{name}")
            key = f"{task_name}:{group}"
            source_group_paths[key] = group_path
            source_group_counts[key] = rows

    source_group_index = json.loads(
        (package / "reports" / "source_label_groups" / "index.json").read_text(encoding="utf-8")
    )
    index_paths = {str(entry["path"]) for entry in source_group_index["entries"]}
    expected_index_paths = {str(path.relative_to(package)) for path in source_group_paths.values()}
    index_checks = {
        "ontology_version": source_group_index["label_ontology_version"]
        == ontology["ontology_version"],
        "training_not_authorized": source_group_index["training_authorized"] is False,
        "harmonized_label_required": source_group_index["unified_training_label_required"]
        == "harmonized_editing_efficiency_fraction",
        "group_file_inventory_exact": index_paths == expected_index_paths,
    }
    for name, passed in index_checks.items():
        checks.append({"check": f"source_label_group_index:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"source_label_group_index_failed:{name}")

    f0 = feature_columns(registry, "F0_core_v1")
    f1 = feature_columns(registry, "F1_sequence_v1")
    context = feature_columns(registry, "F1_context_v1")
    _check_feature_file(package, "features/f0_core_v1.parquet", f0, {"design_id"})
    _check_feature_file(package, "features/f1_sequence_v1.parquet", f1, {"design_id"})
    _check_feature_file(
        package,
        "features/f1_context_v1.parquet",
        context,
        {"observation_id", "design_id"},
    )
    never = {value.lower() for value in registry["never_features"]}
    for feature_set, columns in (("F0", f0), ("F1_sequence", f1), ("F1_context", context)):
        forbidden = sorted(column for column in columns if column.lower() in never)
        passed = not forbidden
        checks.append({"check": f"forbidden_features:{feature_set}", "passed": passed})
        if forbidden:
            raise ValueError(f"forbidden_features:{feature_set}:{forbidden}")
        forbidden_tokens = (
            "source",
            "publication",
            "sheet",
            "excel_row",
            "outcome",
            "label",
            "prediction",
            "score",
            "upstream",
            "split",
            "fold",
            "model",
        )
        leaked = sorted(
            column for column in columns if set(column.lower().split("_")) & set(forbidden_tokens)
        )
        passed = not leaked
        checks.append({"check": f"forbidden_feature_tokens:{feature_set}", "passed": passed})
        if leaked:
            raise ValueError(f"forbidden_feature_tokens:{feature_set}:{leaked}")

    f0_frame = pd.read_parquet(package / "features" / "f0_core_v1.parquet")
    designs = pd.read_parquet(package / "tables" / "designs.parquet")
    target_loci_path = package / "tables" / "target_loci.parquet"
    target_loci = pd.read_parquet(target_loci_path) if target_loci_path.exists() else None
    design_features = designs[["design_id", "pbs", "rtt", "extension"]].merge(
        f0_frame,
        on="design_id",
        how="outer",
        validate="one_to_one",
        indicator=True,
    )
    invariants = {
        "extension_length": bool(
            (f0_frame["extension_length"] == f0_frame["pbs_length"] + f0_frame["rtt_length"]).all()
        ),
        "edit_one_hot": bool(
            (
                f0_frame[
                    ["edit_type_substitution", "edit_type_insertion", "edit_type_deletion"]
                ].sum(axis=1)
                == 1
            ).all()
        ),
        "nonnegative_rha": bool((f0_frame["right_homology_arm_length"] >= 0).all()),
        "design_feature_row_conservation": bool(design_features["_merge"].eq("both").all()),
        "pbs_length_matches_sequence": bool(
            design_features["pbs_length"].eq(design_features["pbs"].str.len()).all()
        ),
        "rtt_length_matches_sequence": bool(
            design_features["rtt_length"].eq(design_features["rtt"].str.len()).all()
        ),
        "extension_matches_rtt_plus_pbs": bool(
            design_features["extension"].eq(design_features["rtt"] + design_features["pbs"]).all()
        ),
        "edit_position_in_rtt": bool(
            design_features["edit_position"].ge(1).all()
            and design_features["edit_position"].le(design_features["rtt_length"]).all()
        ),
    }
    for name, passed in invariants.items():
        checks.append({"check": f"f0_invariant:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"f0_invariant_failed:{name}")

    if target_loci is not None:
        target_locus_checks = {
            "unique_target_locus_id": not target_loci["target_locus_id"].duplicated().any(),
            "window_length_74": bool(target_loci["oriented_74nt_window"].str.len().eq(74).all()),
            "window_dna_alphabet": bool(
                target_loci["oriented_74nt_window"].str.fullmatch("[ACGT]{74}").all()
            ),
            "id_matches_window_hash": bool(
                target_loci.apply(
                    lambda row: (
                        row["target_locus_id"]
                        == stable_hash(
                            "target",
                            {"oriented_74nt_window": row["oriented_74nt_window"]},
                        )
                    ),
                    axis=1,
                ).all()
            ),
        }
        for name, passed in target_locus_checks.items():
            checks.append({"check": f"target_locus_invariant:{name}", "passed": passed})
            if not passed:
                raise ValueError(f"target_locus_invariant_failed:{name}")

    f1_frame = pd.read_parquet(package / "features" / "f1_sequence_v1.parquet")
    for region in ("spacer", "pbs", "rtt", "extension"):
        fraction_columns = [f"base_fraction_{base}_{region}" for base in "acgt"]
        passed = bool(f1_frame[fraction_columns].sum(axis=1).sub(1.0).abs().le(1e-12).all())
        checks.append({"check": f"f1_invariant:base_fraction_sum:{region}", "passed": passed})
        if not passed:
            raise ValueError(f"f1_base_fraction_sum_failed:{region}")

    f1_with_edit = designs[["design_id", "edit_type"]].merge(
        f1_frame[["design_id", "tm_rtt_dna_dna_mismatch"]],
        on="design_id",
        validate="one_to_one",
    )
    indel = f1_with_edit["edit_type"].isin(["insertion", "deletion"])
    substitution = f1_with_edit["edit_type"].eq("substitution")
    mismatch_checks = {
        "indel_mismatch_tm_is_null": bool(
            f1_with_edit.loc[indel, "tm_rtt_dna_dna_mismatch"].isna().all()
        ),
        "substitution_mismatch_tm_has_supported_values": bool(
            f1_with_edit.loc[substitution, "tm_rtt_dna_dna_mismatch"].notna().any()
        ),
    }
    for name, passed in mismatch_checks.items():
        checks.append({"check": f"f1_invariant:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"f1_invariant_failed:{name}")

    on_target = pd.read_parquet(
        package / "tasks" / "pegRNA_on_target.parquet",
        columns=[
            "observation_id",
            "source_record_id",
            "design_id",
            "target_locus_id",
            "source_study",
            "source_sheet",
            "outcome_source_column",
            "editing_efficiency_fraction",
            "label_ontology_version",
            "label_definition_id",
            "label_compatibility_group",
            "label_evidence_status",
            "normalization_status",
            "failure_stage",
            "failure_code",
        ],
    )
    context_rows = pq.ParquetFile(package / "features" / "f1_context_v1.parquet").metadata.num_rows
    observation_checks = {
        "unique_observation_id": not on_target["observation_id"].duplicated().any(),
        "context_row_conservation": context_rows == len(on_target),
        "absolute_fraction_range": bool(
            on_target["editing_efficiency_fraction"].dropna().between(0, 1).all()
        ),
        "quarantined_rows_have_failure_code": bool(
            on_target.loc[
                on_target["normalization_status"].eq("quarantined"),
                ["failure_stage", "failure_code"],
            ]
            .notna()
            .all(axis=None)
        ),
    }
    for name, passed in observation_checks.items():
        checks.append({"check": f"observation_invariant:{name}", "passed": bool(passed)})
        if not passed:
            raise ValueError(f"observation_invariant_failed:{name}")

    chromatin_labels = pd.read_parquet(
        package / "tasks" / "chromatin_context.parquet",
        columns=[
            "observation_id",
            "source_study",
            "source_sheet",
            "outcome_source_column",
            "editing_efficiency_fraction",
            "label_ontology_version",
            "label_definition_id",
            "label_compatibility_group",
            "label_evidence_status",
            "normalization_status",
        ],
    )
    mapping_columns = [
        "source_study",
        "source_sheet",
        "outcome_source_column",
        "label_ontology_version",
        "label_definition_id",
        "label_compatibility_group",
        "label_evidence_status",
    ]
    observed_mapping_rows = pd.concat(
        [on_target[mapping_columns], chromatin_labels[mapping_columns]], ignore_index=True
    ).drop_duplicates()
    expected_mapping_rows: list[tuple[str, ...]] = []
    for mapping in ontology["source_mappings"]:
        definition_id = str(mapping["label_definition_id"])
        definition = ontology["definitions"][definition_id]
        expected_mapping_rows.append(
            (
                str(mapping["source_study"]),
                str(mapping["source_sheet"]),
                str(mapping["outcome_source_column"]),
                str(ontology["ontology_version"]),
                definition_id,
                str(definition["label_compatibility_group"]),
                str(definition["evidence_status"]),
            )
        )
    observed_mapping_set = {
        tuple(str(value) for value in row)
        for row in observed_mapping_rows.itertuples(index=False, name=None)
    }
    expected_mapping_set = set(expected_mapping_rows)
    ontology_checks = {
        "expected_mapping_count": len(expected_mapping_set)
        == int(ontology["expected_source_mapping_count"]),
        "task_mapping_exact": observed_mapping_set == expected_mapping_set,
        "on_target_registration_complete": bool(
            on_target[mapping_columns[3:]].notna().all(axis=None)
        ),
        "chromatin_registration_complete": bool(
            chromatin_labels[mapping_columns[3:]].notna().all(axis=None)
        ),
    }
    for name, passed in ontology_checks.items():
        checks.append({"check": f"label_ontology_invariant:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"label_ontology_invariant_failed:{name}")

    expected_clean_ids = set(
        on_target.loc[
            on_target["normalization_status"].eq("ok")
            & on_target["editing_efficiency_fraction"].notna()
            & on_target["design_id"].notna()
            & on_target["target_locus_id"].notna(),
            "observation_id",
        ]
    )
    mixed = pd.read_parquet(
        mixed_path,
        columns=[
            "observation_id",
            "design_id",
            "target_locus_id",
            "editing_efficiency_fraction",
            "label_compatibility_group",
            "normalization_status",
        ],
    )
    mixed_view_checks = {
        "unique_observation_id": not mixed["observation_id"].duplicated().any(),
        "all_status_ok": bool(mixed["normalization_status"].eq("ok").all()),
        "required_fields_non_null": bool(
            mixed[
                [
                    "design_id",
                    "target_locus_id",
                    "editing_efficiency_fraction",
                    "label_compatibility_group",
                ]
            ]
            .notna()
            .all(axis=None)
        ),
        "exact_mandatory_filter": set(mixed["observation_id"]) == expected_clean_ids,
        "contains_multiple_incompatible_groups": mixed["label_compatibility_group"].nunique() > 1,
    }
    for name, passed in mixed_view_checks.items():
        checks.append({"check": f"mixed_label_audit_invariant:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"mixed_label_audit_invariant_failed:{name}")

    on_target_group_ids: set[str] = set()
    on_target_group_row_total = 0
    for key, path in source_group_paths.items():
        task_name, group = key.split(":", 1)
        group_frame = pd.read_parquet(
            path,
            columns=[
                "observation_id",
                "editing_efficiency_fraction",
                "label_compatibility_group",
                "normalization_status",
            ],
        )
        group_checks = {
            "unique_observation_id": not group_frame["observation_id"].duplicated().any(),
            "single_declared_group": bool(group_frame["label_compatibility_group"].eq(group).all()),
            "all_status_ok": bool(group_frame["normalization_status"].eq("ok").all()),
            "label_non_null": bool(group_frame["editing_efficiency_fraction"].notna().all()),
        }
        for name, passed in group_checks.items():
            checks.append(
                {
                    "check": f"source_label_group_invariant:{task_name}:{group}:{name}",
                    "passed": passed,
                }
            )
            if not passed:
                raise ValueError(f"source_label_group_invariant_failed:{task_name}:{group}:{name}")
        ids = set(group_frame["observation_id"])
        if task_name == "pegRNA_on_target":
            if on_target_group_ids & ids:
                raise ValueError(f"source_label_group_overlap:{group}")
            on_target_group_ids.update(ids)
            on_target_group_row_total += len(group_frame)

    group_union_checks = {
        "on_target_union_exact": on_target_group_ids == expected_clean_ids,
        "on_target_row_conservation": on_target_group_row_total == len(expected_clean_ids),
    }
    expected_chromatin_ids = set(
        chromatin_labels.loc[
            chromatin_labels["normalization_status"].eq("ok")
            & chromatin_labels["editing_efficiency_fraction"].notna(),
            "observation_id",
        ]
    )
    chromatin_group_ids: set[str] = set()
    chromatin_group_total = 0
    for key, path in source_group_paths.items():
        task_name, _group = key.split(":", 1)
        if task_name != "chromatin_context":
            continue
        ids = set(pd.read_parquet(path, columns=["observation_id"])["observation_id"])
        if chromatin_group_ids & ids:
            raise ValueError("chromatin_source_label_group_overlap")
        chromatin_group_ids.update(ids)
        chromatin_group_total += len(ids)
    group_union_checks.update(
        {
            "chromatin_union_exact": chromatin_group_ids == expected_chromatin_ids,
            "chromatin_row_conservation": chromatin_group_total == len(expected_chromatin_ids),
        }
    )
    for name, passed in group_union_checks.items():
        checks.append({"check": f"source_label_group_union:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"source_label_group_union_failed:{name}")

    if target_loci is not None:
        on_target_loci = set(on_target["target_locus_id"].dropna())
        target_locus_id_set = set(target_loci["target_locus_id"])
        passed = on_target_loci <= target_locus_id_set
        checks.append(
            {
                "check": "target_locus_invariant:on_target_ids_resolve",
                "passed": passed,
            }
        )
        if not passed:
            raise ValueError("target_locus_invariant_failed:on_target_ids_resolve")

    recoveries_path = package / "reports" / "deterministic_recoveries.parquet"
    if recoveries_path.exists():
        recoveries = pd.read_parquet(recoveries_path)
        expected_recoveries = int(
            contract["inputs"]["pridict2"]["expected_rows"][
                "sheet_2_deterministic_target_pair_recoveries"
            ]
        )
        recovery_checks = {
            "expected_rows": len(recoveries) == expected_recoveries,
            "unique_source_record_id": not recoveries["source_record_id"].duplicated().any(),
            "registered_code": set(recoveries["recovery_code"])
            == {"wide_target_recovered_from_train_pair"},
        }
        for name, passed in recovery_checks.items():
            checks.append({"check": f"recovery_invariant:{name}", "passed": passed})
            if not passed:
                raise ValueError(f"recovery_invariant_failed:{name}")

    if "mismatch_missingness" in contract:
        mismatch_frame = pd.read_parquet(
            package / "tasks" / "mismatch_off_target.parquet",
            columns=[
                "design_id",
                "target_locus_id",
                "source_group",
                "source_target_group_id",
                "linkage_status",
                "linkage_failure_code",
                "source_study",
                "target_number",
                "normalization_status",
                "failure_code",
                "mismatch_count",
            ],
        )
        source_unknown = mismatch_frame["failure_code"].eq(
            "source_declared_unknown_mismatch_metadata"
        )
        expected_source_unknown = int(
            contract["mismatch_missingness"]["expected_source_unknown_rows"]
        )
        source_unknown_checks = {
            "expected_rows": int(source_unknown.sum()) == expected_source_unknown,
            "status": bool(
                mismatch_frame.loc[source_unknown, "normalization_status"]
                .eq("complete_with_source_unknown_fields")
                .all()
            ),
            "mismatch_count_is_null": bool(
                mismatch_frame.loc[source_unknown, "mismatch_count"].isna().all()
            ),
        }
        for name, passed in source_unknown_checks.items():
            checks.append(
                {
                    "check": f"source_unknown_mismatch_invariant:{name}",
                    "passed": passed,
                }
            )
            if not passed:
                raise ValueError(f"source_unknown_mismatch_invariant_failed:{name}")

        expected_groups = int(contract["mismatch_linkage"]["expected_source_target_groups"])
        linkage_checks = {
            "canonical_design_id_explicitly_null": bool(mismatch_frame["design_id"].isna().all()),
            "canonical_target_locus_id_explicitly_null": bool(
                mismatch_frame["target_locus_id"].isna().all()
            ),
            "source_group_complete": bool(mismatch_frame["source_group"].notna().all()),
            "source_target_group_id_complete": bool(
                mismatch_frame["source_target_group_id"].notna().all()
            ),
            "source_target_group_count": (
                mismatch_frame["source_target_group_id"].nunique() == expected_groups
            ),
            "source_target_group_id_deterministic": bool(
                mismatch_frame.apply(
                    lambda row: (
                        row["source_target_group_id"]
                        == stable_hash(
                            "source_target_group",
                            {
                                "source_study": row["source_study"],
                                "source_group": row["source_group"],
                                "target_number": row["target_number"],
                            },
                        )
                    ),
                    axis=1,
                ).all()
            ),
            "linkage_status_explicit": bool(
                mismatch_frame["linkage_status"].eq("source_group_only").all()
            ),
            "linkage_failure_code_explicit": bool(
                mismatch_frame["linkage_failure_code"]
                .eq("canonical_sequence_fields_absent_from_source_sheet")
                .all()
            ),
        }
        for name, passed in linkage_checks.items():
            checks.append({"check": f"mismatch_linkage_invariant:{name}", "passed": passed})
            if not passed:
                raise ValueError(f"mismatch_linkage_invariant_failed:{name}")

    near_pairs_path = package / "reports" / "cross_source_near_target_pairs.parquet"
    if target_loci is not None and near_pairs_path.exists():
        near_pairs = pd.read_parquet(near_pairs_path)
        target_windows = target_loci.set_index("target_locus_id")["oriented_74nt_window"]
        deepprime_target_ids = set(
            on_target.loc[on_target["source_study"].eq("DeepPrime"), "target_locus_id"]
            .dropna()
            .unique()
        )
        pridict2_target_ids = set(
            on_target.loc[on_target["source_study"].eq("PRIDICT2"), "target_locus_id"]
            .dropna()
            .unique()
        )
        near_pair_checks = {
            "unique_pairs": not near_pairs[
                ["deepprime_target_locus_id", "pridict2_target_locus_id"]
            ]
            .duplicated()
            .any(),
            "cross_source_membership": bool(
                set(near_pairs["deepprime_target_locus_id"]) <= deepprime_target_ids
                and set(near_pairs["pridict2_target_locus_id"]) <= pridict2_target_ids
            ),
            "distance_exactly_one": bool(
                near_pairs.apply(
                    lambda row: (
                        sum(
                            left != right
                            for left, right in zip(
                                target_windows[row["deepprime_target_locus_id"]],
                                target_windows[row["pridict2_target_locus_id"]],
                                strict=True,
                            )
                        )
                        == 1
                    ),
                    axis=1,
                ).all()
            ),
        }
        for name, passed in near_pair_checks.items():
            checks.append({"check": f"near_target_invariant:{name}", "passed": passed})
            if not passed:
                raise ValueError(f"near_target_invariant_failed:{name}")

    source_anomaly_report = json.loads(
        (package / "reports" / "source_anomalies.json").read_text(encoding="utf-8")
    )
    expected_anomaly_ids = {str(item["id"]) for item in contract.get("known_source_anomalies", [])}
    observed_anomaly_ids = {
        str(item["anomaly_id"]) for item in source_anomaly_report["registered_anomalies"]
    }
    source_anomaly_checks = {
        "status_pass": source_anomaly_report["status"] == "pass",
        "registered_ids_exact": observed_anomaly_ids == expected_anomaly_ids,
        "all_registered_checks_pass": all(
            item["passed"] for item in source_anomaly_report["registered_anomalies"]
        ),
        "no_unregistered_anomalies": not source_anomaly_report["unregistered_anomalies"],
    }
    for name, passed in source_anomaly_checks.items():
        checks.append({"check": f"source_anomaly_invariant:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"source_anomaly_invariant_failed:{name}")

    source_recalculation_report = json.loads(
        (package / "reports" / "source_recalculation_audit.json").read_text(encoding="utf-8")
    )
    interpretation = source_recalculation_report.get("interpretation", {})
    passed = bool(interpretation.get("pridict2_base_after_edit"))
    checks.append(
        {
            "check": "source_recalculation:pridict2_base_after_edit_semantics_documented",
            "passed": passed,
        }
    )
    if not passed:
        raise ValueError(
            "source_recalculation_invariant_failed:pridict2_base_after_edit_semantics_documented"
        )

    lineage_ids = pd.read_parquet(
        package / "tables" / "lineage.parquet", columns=["source_record_id"]
    )["source_record_id"]
    raw_ids: list[str] = []
    for raw_path in sorted((package / "raw_mirror").glob("*.parquet")):
        raw_ids.extend(
            pq.read_table(raw_path, columns=["__source_record_id"])
            .column("__source_record_id")
            .to_pylist()
        )
    raw_id_set = set(raw_ids)
    lineage_id_set = set(lineage_ids)
    row_conservation_checks = {
        "raw_and_lineage_rows": len(raw_ids) == len(lineage_ids),
        "raw_ids_unique": len(raw_ids) == len(raw_id_set),
        "lineage_ids_unique": len(lineage_ids) == len(lineage_id_set),
        "raw_and_lineage_ids_equal": raw_id_set == lineage_id_set,
    }
    for name, passed in row_conservation_checks.items():
        checks.append({"check": f"row_conservation:{name}", "passed": passed})
        if not passed:
            raise ValueError(f"row_conservation_failed:{name}")

    task_source_ids = set(on_target["source_record_id"])
    for relative in (
        "tasks/mismatch_off_target.parquet",
        "tasks/chromatin_context.parquet",
        "quarantine/rows.parquet",
    ):
        task_source_ids.update(
            pq.read_table(package / relative, columns=["source_record_id"])
            .column("source_record_id")
            .to_pylist()
        )
    passed = task_source_ids <= lineage_id_set
    checks.append(
        {"check": "row_conservation:normalized_and_quarantine_ids_in_lineage", "passed": passed}
    )
    if not passed:
        raise ValueError("normalized_or_quarantine_id_missing_from_lineage")

    return {
        "status": "pass",
        "package": package_label or str(package),
        "task_counts": task_counts,
        "label_ontology_version": ontology["ontology_version"],
        "label_harmonization_version": harmonization["harmonization_version"],
        "harmonized_label_status": manifest["harmonized_label_status"],
        "source_label_group_audit_counts": source_group_counts,
        "feature_widths": {
            "F0_core_v1": len(f0),
            "F1_sequence_v1": len(f1),
            "F1_context_v1": len(context),
        },
        "checks": checks,
    }
