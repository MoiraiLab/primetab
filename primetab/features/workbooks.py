"""Workbook contracts and source-specific canonicalization."""

from __future__ import annotations

import ast
import json
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.read_only import ReadOnlyCell

from primetab.features.features import (
    CanonicalDesign,
    local_bases,
    normalize_dna,
    parse_edit_type,
    reconstruct_edited_target,
    reverse_complement,
    stable_hash,
    validate_edit_alignment,
)

DEEPPRIME_WIDE_TARGET_PREFIX = "Wide target sequence"
DEEPPRIME_MASKED_TARGET_PREFIX = "Edited target sequence"
PRIDICT2_OPTIMIZED_SCAFFOLD = (
    "GTTTCAGAGCTATGCTGGAAACAGCATAGCAAGTTGAAATAAGGCTAGTCCGTTATCAACTTGAAAAAGTGGCA"
    "CCGAGTCGGTGC"
)


@dataclass(frozen=True)
class SheetSpec:
    header_row: int
    first_data_row: int
    data_rows: int | None = None


def source_record_id(workbook_sha256: str, sheet: str, excel_row: int) -> str:
    return stable_hash(
        "source",
        {"workbook_sha256": workbook_sha256, "sheet": sheet, "excel_row": excel_row},
    )


def condition_id(condition: dict[str, str]) -> str:
    return stable_hash("condition", condition)


def observation_id(record_id: str, outcome_column: str) -> str:
    return stable_hash(
        "observation",
        {"source_record_id": record_id, "outcome_column": outcome_column},
    )


def target_locus_id(window: str) -> str:
    return stable_hash("target", {"oriented_74nt_window": window})


def trim_headers(values: Iterable[Any]) -> list[Any]:
    headers = list(values)
    while headers and headers[-1] is None:
        headers.pop()
    return headers


def unique_headers(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(trim_headers(values), start=1):
        base = str(value).strip() if value is not None else f"unnamed_{index:03d}"
        base = base or f"unnamed_{index:03d}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return result


def raw_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def iter_sheet_records(
    worksheet: Any,
    spec: SheetSpec,
    workbook_sha256: str,
) -> tuple[list[str], Iterator[tuple[int, str, dict[str, Any], list[ReadOnlyCell]]]]:
    header_cells = next(
        worksheet.iter_rows(min_row=spec.header_row, max_row=spec.header_row)
    )
    headers = unique_headers(cell.value for cell in header_cells)
    final_row = (
        spec.first_data_row + spec.data_rows - 1
        if spec.data_rows is not None
        else worksheet.max_row
    )

    def records() -> Iterator[tuple[int, str, dict[str, Any], list[ReadOnlyCell]]]:
        for excel_row, cells in enumerate(
            worksheet.iter_rows(
                min_row=spec.first_data_row,
                max_row=final_row,
                max_col=len(headers),
            ),
            start=spec.first_data_row,
        ):
            cell_list = list(cells)
            values = [cell.value for cell in cell_list]
            if all(value is None for value in values):
                continue
            record = dict(zip(headers, values, strict=True))
            yield (
                excel_row,
                source_record_id(workbook_sha256, worksheet.title, excel_row),
                record,
                cell_list,
            )

    return headers, records()


def load_deepprime_summary(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Summary"]
        headers = unique_headers(
            cell.value for cell in next(worksheet.iter_rows(min_row=5, max_row=5))
        )
        rows = []
        for values in worksheet.iter_rows(min_row=6, max_col=len(headers), values_only=True):
            if values[0] is None:
                continue
            record = dict(zip(headers, values, strict=True))
            record["Index"] = int(record["Index"])
            record["Total dataset size"] = int(record["Total dataset size"])
            rows.append(record)
        return rows
    finally:
        workbook.close()


def deepprime_specs(summary: list[dict[str, Any]]) -> dict[str, SheetSpec]:
    specs = {"Summary": SheetSpec(header_row=5, first_data_row=6, data_rows=len(summary))}
    for row in summary:
        specs[str(row["Index"])] = SheetSpec(
            header_row=4,
            first_data_row=5,
            data_rows=int(row["Total dataset size"]),
        )
    return specs


def pridict2_specs(path: Path) -> dict[str, SheetSpec]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        result: dict[str, SheetSpec] = {
            "Overview": SheetSpec(header_row=4, first_data_row=5)
        }
        for name in workbook.sheetnames:
            if name != "Overview":
                result[name] = SheetSpec(header_row=1, first_data_row=2)
        return result
    finally:
        workbook.close()


def find_column(record: dict[str, Any], prefix: str) -> Any:
    matches = [value for key, value in record.items() if key.startswith(prefix)]
    if len(matches) != 1:
        raise KeyError(f"expected_one_column_with_prefix:{prefix}:{len(matches)}")
    return matches[0]


def as_int(value: Any, name: str) -> int:
    if value is None:
        raise ValueError(f"missing_integer:{name}")
    number = int(float(value))
    if float(value) != number:
        raise ValueError(f"non_integral:{name}:{value}")
    return number


def _deepprime_edit_type(record: dict[str, Any]) -> str:
    flags = {
        "substitution": as_int(record["type_sub"], "type_sub"),
        "insertion": as_int(record["type_ins"], "type_ins"),
        "deletion": as_int(record["type_del"], "type_del"),
    }
    active = [name for name, enabled in flags.items() if enabled == 1]
    if len(active) != 1 or any(value not in {0, 1} for value in flags.values()):
        raise ValueError(f"invalid_edit_type_one_hot:{flags}")
    return active[0]


def canonicalize_deepprime(
    record: dict[str, Any],
    *,
    allow_additional_rtt_mismatches: bool = False,
) -> tuple[CanonicalDesign, str, str]:
    wide_target = normalize_dna(find_column(record, DEEPPRIME_WIDE_TARGET_PREFIX))
    if len(wide_target) != 74:
        raise ValueError(f"wide_target_length:{len(wide_target)}")
    masked = str(find_column(record, DEEPPRIME_MASKED_TARGET_PREFIX)).strip().upper()
    if len(masked) != 74 or set(masked) - set("ACGTX"):
        raise ValueError("invalid_masked_target")

    pbs_length = as_int(record["PBSlen"], "PBSlen")
    rtt_length = as_int(record["RTlen"], "RTlen")
    edit_position = as_int(record["Edit_pos"], "Edit_pos")
    edit_length = as_int(record["Edit_len"], "Edit_len")
    edit_type = _deepprime_edit_type(record)
    start = 21 - pbs_length
    end = 21 + rtt_length
    combined = masked[start:end]
    if len(combined) != pbs_length + rtt_length or "X" in combined:
        raise ValueError("masked_extension_not_recoverable")
    # DeepPrime masks target-oriented PBS+RTT sequence.  The unified contract
    # stores both regions in the 5'-to-3' pegRNA transcript direction.
    pbs_target = normalize_dna(combined[:pbs_length])
    rtt_target = normalize_dna(combined[pbs_length:])
    pbs = reverse_complement(pbs_target)
    rtt = reverse_complement(rtt_target)
    if as_int(record["RT-PBSlen"], "RT-PBSlen") != len(pbs) + len(rtt):
        raise ValueError("source_extension_length_mismatch")

    spacer = wide_target[4:24]
    pam_before = wide_target[24:27]
    edited_target, rtt_original_target = reconstruct_edited_target(
        wide_target, 21, rtt_target, edit_type, edit_length
    )
    validate_edit_alignment(
        rtt_original_target,
        rtt_target,
        edit_type,
        edit_position,
        edit_length,
        allow_additional_mismatches=allow_additional_rtt_mismatches,
    )
    rtt_original = reverse_complement(rtt_original_target)
    pam_after = edited_target[24:27]
    bases = local_bases(wide_target, 21, edit_position, edit_type, edit_length)
    rha = (
        len(rtt) - edit_position + 1
        if edit_type == "deletion"
        else len(rtt) - edit_position - edit_length + 1
    )
    if rha < 0:
        raise ValueError("negative_right_homology_arm")
    design = CanonicalDesign(
        spacer=spacer,
        pbs=pbs,
        rtt=rtt,
        pam_before=pam_before,
        pam_after=pam_after,
        edit_type=edit_type,
        edit_position=edit_position,
        edit_length=edit_length,
        right_homology_arm_length=rha,
        rtt_original=rtt_original,
        **bases,
    )
    return design, target_locus_id(wide_target), wide_target


def parse_location(value: Any, name: str) -> tuple[int, int]:
    try:
        parsed = ast.literal_eval(str(value))
        if not isinstance(parsed, (list, tuple)) or len(parsed) != 2:
            raise ValueError
        return int(parsed[0]), int(parsed[1])
    except (SyntaxError, ValueError, TypeError) as exc:
        raise ValueError(f"invalid_location:{name}:{value}") from exc


def _locate_canonical_spacer(wide_target: str, peg_spacer: str) -> tuple[str, int]:
    exact = wide_target.find(peg_spacer)
    if exact >= 0:
        return peg_spacer, exact
    suffix = wide_target.find(peg_spacer[1:])
    if suffix > 0:
        start = suffix - 1
        return wide_target[start : start + 20], start
    raise ValueError("spacer_not_found_in_target")


def _target_74_window(wide_target: str, spacer_start: int) -> str:
    start = spacer_start - 4
    end = start + 74
    if start < 0 or end > len(wide_target):
        raise ValueError("target_74_window_out_of_bounds")
    return wide_target[start:end]


def pridict2_sheet2_wide_sequences(
    record: dict[str, Any],
) -> tuple[str, str, str | None]:
    """Return the source target pair, using only an exact redundant-field fallback.

    Thirty-six label-free deletion designs omit the two ``wide_*_target`` cells but
    retain the corresponding 99-nt ``wide_*_train`` pair.  The caller still applies
    all spacer, RTT, edit-alignment, and 74-nt-window checks, so this fallback does
    not guess or impute sequence content.
    """

    target_initial = record.get("wide_initial_target")
    target_mutated = record.get("wide_mutated_target")
    if target_initial is not None and target_mutated is not None:
        return normalize_dna(target_initial), normalize_dna(target_mutated), None
    if target_initial is not None or target_mutated is not None:
        raise ValueError("incomplete_wide_target_pair")

    train_initial = record.get("wide_initial_train")
    train_mutated = record.get("wide_mutated_train")
    if train_initial is None or train_mutated is None:
        raise ValueError("missing_sequence")
    return (
        normalize_dna(train_initial),
        normalize_dna(train_mutated),
        "wide_target_recovered_from_train_pair",
    )


def canonicalize_pridict2_sheet2(
    record: dict[str, Any],
) -> tuple[CanonicalDesign, str, str]:
    wide_initial, wide_mutated, _recovery_code = pridict2_sheet2_wide_sequences(
        record
    )
    peg_spacer = normalize_dna(record["spacer"])
    spacer, spacer_start = _locate_canonical_spacer(wide_initial, peg_spacer)
    pbs = normalize_dna(record["PBS"])
    rtt = normalize_dna(record["RTT"])
    extension = normalize_dna(record["Extension_seq"])
    if extension != rtt + pbs:
        raise ValueError("extension_not_rtt_plus_pbs")
    if as_int(record["PBSlength"], "PBSlength") != len(pbs):
        raise ValueError("pbs_length_mismatch")
    if as_int(record["RTTlength"], "RTTlength") != len(rtt):
        raise ValueError("rtt_length_mismatch")

    edit_type = parse_edit_type(record["Correction_Type"])
    edit_length = as_int(record["Correction_Length"], "Correction_Length")
    # Supplementary Table 2 uses a zero-based edit position, whereas v1 uses
    # the first target base after the primary nick as position 1.
    edit_position = as_int(record["Editing_Position"], "Editing_Position") + 1
    pam_start = spacer_start + 20
    pam_before = wide_initial[pam_start : pam_start + 3]
    pam_after = wide_mutated[pam_start : pam_start + 3]
    if len(pam_before) != 3 or len(pam_after) != 3:
        raise ValueError("pam_out_of_bounds")

    rt_start, rt_end = parse_location(record["RT_initial_location"], "RT_initial_location")
    rtt_original_target = wide_initial[rt_start:rt_end]
    expected_original_length = len(rtt) + (
        edit_length if edit_type == "deletion" else -edit_length if edit_type == "insertion" else 0
    )
    if len(rtt_original_target) != expected_original_length:
        raise ValueError("rtt_original_length_mismatch")
    mutated_start, mutated_end = parse_location(
        record["RT_mutated_location"], "RT_mutated_location"
    )
    rtt_edited_target = wide_mutated[mutated_start:mutated_end]
    if rtt_edited_target != reverse_complement(rtt):
        raise ValueError("rtt_mutated_sequence_mismatch")
    validate_edit_alignment(
        rtt_original_target,
        rtt_edited_target,
        edit_type,
        edit_position,
        edit_length,
    )
    rtt_original = reverse_complement(rtt_original_target)
    nick_index = rt_start
    bases = local_bases(wide_initial, nick_index, edit_position, edit_type, edit_length)
    rha = (
        len(rtt) - edit_position + 1
        if edit_type == "deletion"
        else len(rtt) - edit_position - edit_length + 1
    )
    if rha < 0:
        raise ValueError("negative_right_homology_arm")
    if (
        "RTToverhanglength" in record
        and record["RTToverhanglength"] is not None
        and as_int(record["RTToverhanglength"], "RTToverhanglength") != rha
    ):
        raise ValueError("source_rha_mismatch")
    design = CanonicalDesign(
        spacer=spacer,
        pbs=pbs,
        rtt=rtt,
        pam_before=pam_before,
        pam_after=pam_after,
        edit_type=edit_type,
        edit_position=edit_position,
        edit_length=edit_length,
        right_homology_arm_length=rha,
        rtt_original=rtt_original,
        **bases,
    )
    window = _target_74_window(wide_initial, spacer_start)
    return design, target_locus_id(window), window


def canonicalize_pridict2_sheet7(
    record: dict[str, Any],
) -> tuple[CanonicalDesign, str, str]:
    peg_rna = normalize_dna(record["pegRNA_sequence"])
    scaffold_start = peg_rna.find(PRIDICT2_OPTIMIZED_SCAFFOLD)
    if scaffold_start <= 0:
        raise ValueError("optimized_scaffold_not_found")
    peg_spacer = peg_rna[:scaffold_start]
    if len(peg_spacer) != 20:
        raise ValueError(f"unexpected_peg_spacer_length:{len(peg_spacer)}")
    extension = peg_rna[scaffold_start + len(PRIDICT2_OPTIMIZED_SCAFFOLD) :]
    pbs_length = as_int(record["PBSlength"], "PBSlength")
    rtt_length = as_int(record["RTTlength"], "RTTlength")
    if len(extension) != pbs_length + rtt_length:
        raise ValueError("arrayed_extension_length_mismatch")
    rtt = extension[:rtt_length]
    pbs = extension[rtt_length:]

    wide_initial = normalize_dna(record["original_target"])
    spacer, spacer_start = _locate_canonical_spacer(wide_initial, peg_spacer)
    edit_type = parse_edit_type(record["Correction_Type"])
    edit_length = as_int(record["Edit_length"], "Edit_length")
    edit_position = as_int(record["Edit_position"], "Edit_position")
    nick_index = spacer_start + 17
    original_length = len(rtt) + (
        edit_length if edit_type == "deletion" else -edit_length if edit_type == "insertion" else 0
    )
    if original_length <= 0 or nick_index + original_length > len(wide_initial):
        raise ValueError("arrayed_rtt_original_out_of_bounds")
    rtt_original_target = wide_initial[nick_index : nick_index + original_length]
    rtt_original = reverse_complement(rtt_original_target)
    rtt_edited_target = reverse_complement(rtt)
    validate_edit_alignment(
        rtt_original_target,
        rtt_edited_target,
        edit_type,
        edit_position,
        edit_length,
    )
    edited_target = (
        wide_initial[:nick_index]
        + rtt_edited_target
        + wide_initial[nick_index + original_length :]
    )
    pam_start = spacer_start + 20
    pam_before = wide_initial[pam_start : pam_start + 3]
    pam_after = edited_target[pam_start : pam_start + 3]
    bases = local_bases(wide_initial, nick_index, edit_position, edit_type, edit_length)
    rha = (
        len(rtt) - edit_position + 1
        if edit_type == "deletion"
        else len(rtt) - edit_position - edit_length + 1
    )
    if rha < 0:
        raise ValueError("negative_right_homology_arm")
    if (
        "RTToverhanglength" in record
        and record["RTToverhanglength"] is not None
        and as_int(record["RTToverhanglength"], "RTToverhanglength") != rha
    ):
        raise ValueError("source_rha_mismatch")
    design = CanonicalDesign(
        spacer=spacer,
        pbs=pbs,
        rtt=rtt,
        pam_before=pam_before,
        pam_after=pam_after,
        edit_type=edit_type,
        edit_position=edit_position,
        edit_length=edit_length,
        right_homology_arm_length=rha,
        rtt_original=rtt_original,
        **bases,
    )
    window = _target_74_window(wide_initial, spacer_start)
    return design, target_locus_id(window), window


def deepprime_condition(summary_row: dict[str, Any]) -> dict[str, str]:
    cell = str(summary_row["Cell line"]).strip()
    editor = str(summary_row["PE system"]).strip()
    scaffold = str(summary_row["Scaffold"]).strip().lower() or "unknown"
    organism = "mouse" if cell == "NIH3T3" else "human"
    return {
        "cell_line": cell,
        "editor_system": editor,
        "pam_family": "NRCH" if "NRCH" in editor.upper() else "NGG",
        "scaffold_class": scaffold,
        "organism": organism,
        "mmr_status": "inhibited" if "PE4" in editor.upper() else "unknown",
        "assay_type": "unknown",
        "target_context": "unknown",
        "delivery_context": "VLP" if "VLP" in editor.upper() else "unknown",
    }


def pridict2_sheet2_conditions() -> dict[str, dict[str, str]]:
    base = {
        "editor_system": "PE2",
        "pam_family": "NGG",
        "scaffold_class": "optimized",
        "assay_type": "integrated_reporter",
        "target_context": "integrated",
    }
    return {
        "HEKaverageedited": {
            **base,
            "cell_line": "HEK293T",
            "organism": "human",
            "mmr_status": "deficient",
            "delivery_context": "plasmid_transfection",
        },
        "K562averageedited": {
            **base,
            "cell_line": "K562",
            "organism": "human",
            "mmr_status": "proficient",
            "delivery_context": "plasmid_electroporation",
        },
        "K562MLH1dnaverageedited": {
            **base,
            "cell_line": "K562-MLH1dn",
            "editor_system": "PE4",
            "organism": "human",
            "mmr_status": "inhibited",
            "delivery_context": "plasmid_electroporation",
        },
        "AdVaverageedited": {
            **base,
            "cell_line": "mouse_hepatocyte_in_vivo",
            "editor_system": "PE2deltaRnH",
            "organism": "mouse",
            "mmr_status": "proficient",
            "delivery_context": "AAV_Sleeping_Beauty_library_plus_AdV5_editor",
            "assay_type": "in_vivo",
            "target_context": "mixed_integrated_and_episomal_reporter",
        },
    }


def pridict2_sheet7_conditions() -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    mmr_by_cell = {
        "K562": "proficient",
        "HEK293T": "deficient",
        "HepG2": "unknown",
    }
    for cell in ("K562", "HEK293T", "HepG2"):
        for context in ("endogenous", "reporter"):
            column = f"{cell}_edited_percentage_{context}"
            result[column] = {
                "cell_line": cell,
                "editor_system": "PE2",
                "pam_family": "NGG",
                "scaffold_class": "optimized",
                "organism": "human",
                "mmr_status": mmr_by_cell[cell],
                "assay_type": "endogenous_locus" if context == "endogenous" else "integrated_reporter",
                "target_context": context if context == "endogenous" else "integrated",
                "delivery_context": "unknown",
            }
    return result
